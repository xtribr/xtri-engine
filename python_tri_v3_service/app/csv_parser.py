"""
CSV/XLSX parser for ENEM binary response matrices.

Handles:
- CSV with comma or semicolon delimiters
- XLSX files (first sheet)
- Numeric headers (1, 2, ..., 135) or "Questão N" headers
- Metadata columns skipped: ID, Idioma, Acerto, TRI, CALI, etc.
- User explicitly selects which area the file belongs to
"""

import csv
import io
import re
from typing import Any


# Columns to always skip (case-insensitive match)
_SKIP_HEADERS = {
    "id", "idioma", "acerto", "acertos",
    "tri", "tri 1", "tri 2", "tri1", "tri2", "tri 3",
    "cali1", "cali2", "cali 1", "cali 2",
    "nome", "name", "turma",
}

# Pattern for "Questão N" or "Questao N"
_QUESTAO_RE = re.compile(r"quest[ãa]o\s+(\d+)", re.IGNORECASE)


def _detect_delimiter(sample: str) -> str:
    """Detect delimiter from the first few lines: tab, semicolon, or comma."""
    tabs = sample.count("\t")
    semicolons = sample.count(";")
    commas = sample.count(",")
    if tabs > semicolons and tabs > commas:
        return "\t"
    if semicolons > commas:
        return ";"
    return ","


def _is_question_column(header: str) -> bool:
    """Check if a column header represents a question (binary response)."""
    header = header.strip()
    if not header:
        return False
    if header.lower() in _SKIP_HEADERS:
        return False
    # Direct integer
    try:
        int(header)
        return True
    except ValueError:
        pass
    # "Questão N" pattern
    if _QUESTAO_RE.match(header):
        return True
    return False


def parse_single_area(content: str, is_xlsx: bool = False) -> dict[str, Any]:
    """
    Parse a file containing binary responses for ONE area.

    The user explicitly selects which area this file belongs to.
    The parser only extracts student IDs and binary response vectors.

    Args:
        content: Raw CSV text (decoded), or raw bytes for XLSX
        is_xlsx: If True, content is bytes from an XLSX file

    Returns:
        {
            "alunos": [{"aluno_id": str, "respostas": [0,1,...]}],
            "n_alunos": int,
            "n_itens": int,
            "avisos": ["warning messages..."]
        }
    """
    if is_xlsx:
        return _parse_xlsx_bytes(content)
    return _parse_csv_text(content)


def _is_binary_row(row: list[str]) -> bool:
    """Check if a row is composed entirely of 0/1 values (no header)."""
    non_empty = [c.strip() for c in row if c.strip()]
    if len(non_empty) < 30:
        return False
    return all(c in ("0", "1") for c in non_empty)


def _parse_csv_text(csv_content: str) -> dict[str, Any]:
    """Parse CSV text content."""
    avisos: list[str] = []

    delimiter = _detect_delimiter(csv_content[:2000])
    reader = csv.reader(io.StringIO(csv_content), delimiter=delimiter)

    first_row = next(reader)

    # Headerless detection: if first row is all 0/1 binary, synthesize headers
    headerless = _is_binary_row(first_row)
    if headerless:
        n_cols = len(first_row)
        headers = [str(i + 1) for i in range(n_cols)]
        # Re-create reader to include the first row as data
        reader = csv.reader(io.StringIO(csv_content), delimiter=delimiter)
    else:
        headers = first_row

    # Find ID column and question columns
    id_col: int | None = None
    question_cols: list[int] = []

    for idx, h in enumerate(headers):
        h_clean = h.strip()
        if h_clean.lower() == "id":
            id_col = idx
        elif _is_question_column(h_clean):
            question_cols.append(idx)

    # If no explicit ID column found but first column header is blank/non-question,
    # treat it as the ID column (common format: ";136;137;..." with data "ALUNO1;0;1;...")
    if id_col is None and headers and not headerless and not _is_question_column(headers[0].strip()):
        id_col = 0

    if not question_cols:
        raise ValueError(
            f"No question columns found. "
            f"Headers: {[h.strip() for h in headers[:10]]}..."
        )

    # Parse student rows
    alunos: list[dict] = []
    for row_idx, row in enumerate(reader):
        if not row or all(c.strip() == "" for c in row):
            continue

        aluno_id = (
            row[id_col].strip()
            if id_col is not None and id_col < len(row)
            else f"aluno_{row_idx + 1}"
        )

        respostas: list[int] = []
        for col_idx in question_cols:
            if col_idx < len(row):
                val = row[col_idx].strip()
                if val in ("0", "1"):
                    respostas.append(int(val))
                else:
                    respostas.append(0)
                    if val and val not in ("", " "):
                        avisos.append(
                            f"Row {row_idx + 2}, col {col_idx + 1}: "
                            f"non-binary value '{val}' treated as 0"
                        )
            else:
                respostas.append(0)

        alunos.append({"aluno_id": aluno_id, "respostas": respostas})

    if len(avisos) > 20:
        total = len(avisos)
        avisos = avisos[:20]
        avisos.append(f"... and {total - 20} more warnings")

    n_itens = len(question_cols)
    return {
        "alunos": alunos,
        "n_alunos": len(alunos),
        "n_itens": n_itens,
        "avisos": avisos,
    }


def _parse_xlsx_bytes(raw_bytes: bytes) -> dict[str, Any]:
    """Parse XLSX file bytes (first sheet)."""
    import openpyxl

    wb = openpyxl.load_workbook(
        io.BytesIO(raw_bytes) if isinstance(raw_bytes, bytes) else io.BytesIO(raw_bytes.encode()),
        data_only=True,
    )
    ws = wb[wb.sheetnames[0]]

    headers = [str(c.value) if c.value is not None else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]

    # Find columns
    id_col: int | None = None
    question_cols: list[int] = []

    for idx, h in enumerate(headers):
        h_clean = h.strip()
        if h_clean.lower() == "id":
            id_col = idx
        elif _is_question_column(h_clean):
            question_cols.append(idx)

    if not question_cols:
        raise ValueError(
            f"No question columns found. "
            f"Headers: {headers[:10]}..."
        )

    avisos: list[str] = []
    alunos: list[dict] = []

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True)):
        if all(v is None for v in row):
            continue

        aluno_id = (
            str(row[id_col]).strip()
            if id_col is not None and id_col < len(row) and row[id_col] is not None
            else f"aluno_{row_idx + 1}"
        )

        respostas: list[int] = []
        for col_idx in question_cols:
            if col_idx < len(row) and row[col_idx] is not None:
                val = row[col_idx]
                if val in (0, 1, 0.0, 1.0):
                    respostas.append(int(val))
                elif str(val).strip() in ("0", "1"):
                    respostas.append(int(str(val).strip()))
                else:
                    respostas.append(0)
                    avisos.append(
                        f"Row {row_idx + 2}, col {col_idx + 1}: "
                        f"non-binary value '{val}' treated as 0"
                    )
            else:
                respostas.append(0)

        alunos.append({"aluno_id": aluno_id, "respostas": respostas})

    if len(avisos) > 20:
        total = len(avisos)
        avisos = avisos[:20]
        avisos.append(f"... and {total - 20} more warnings")

    return {
        "alunos": alunos,
        "n_alunos": len(alunos),
        "n_itens": len(question_cols),
        "avisos": avisos,
    }
