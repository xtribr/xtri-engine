"""Endpoints de exportação CSV/Excel (motor principal e validator V3 x V4)."""

import csv
import io
import logging
from pathlib import Path
from typing import Any, Dict

from fastapi import APIRouter, Body, HTTPException
from fastapi.responses import JSONResponse, StreamingResponse

logger = logging.getLogger(__name__)

router = APIRouter()


def _downloads_export_root() -> Path:
    """Diretório visível ao usuário para exportações geradas localmente."""

    export_root = Path.home() / "Downloads" / "XTRI Engine"
    export_root.mkdir(parents=True, exist_ok=True)
    return export_root


def _save_local_export(filename: str, payload: bytes) -> dict[str, Any]:
    """Salva um arquivo de exportação em Downloads e retorna o metadata básico."""

    destination = _downloads_export_root() / filename
    destination.write_bytes(payload)
    return {
        "status": "sucesso",
        "nome_arquivo": destination.name,
        "caminho": str(destination),
        "bytes": len(payload),
    }


def _build_resultados_csv_bytes(payload: Dict[str, Any]) -> tuple[bytes, str]:
    """Gera o CSV do motor principal com BOM para abrir corretamente no Excel."""

    from datetime import datetime

    area = str(payload.get("area", "")).upper()
    ano = int(payload.get("ano_referencia", 0) or 0)
    resultados = payload.get("resultados", [])
    if not area or not isinstance(resultados, list):
        raise ValueError("Payload inválido para exportação CSV")

    area_lower = area.lower()
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";")
    writer.writerow(["ID", "Acertos", "Theta", "Nota_TRI", "Pct_Dificeis", "Nivel"])
    for row in resultados:
        writer.writerow(
            [
                row.get("aluno_id", ""),
                row.get(f"acertos_{area_lower}", ""),
                row.get(f"theta_{area_lower}", ""),
                row.get(f"nota_{area_lower}", ""),
                row.get(f"acert_dif_{area_lower}", ""),
                row.get(f"estado_{area_lower}", ""),
            ]
        )

    csv_bytes = ("\ufeff" + buffer.getvalue()).encode("utf-8")
    filename = f"xtri_{area}_{ano}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    return csv_bytes, filename


def _build_resultados_excel_bytes(payload: Dict[str, Any]) -> tuple[bytes, str]:
    """Gera o XLSX do motor principal para streaming ou save local."""

    import base64
    from datetime import datetime

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    area = str(payload.get("area", "") or "").upper()
    if not area:
        areas_processadas = payload.get("areas_processadas") or []
        if isinstance(areas_processadas, list) and areas_processadas:
            area = str(areas_processadas[0]).upper()
    ano = payload.get("ano_referencia", 0)
    metodologia = payload.get("metodologia") if isinstance(payload.get("metodologia"), dict) else {}
    metodo = payload.get("metodo") or metodologia.get("metodo_principal", "")
    n_alunos = payload.get("n_alunos", 0)
    resultados = payload.get("resultados", [])
    referencia = payload.get("referencia", [])
    diagnosticos = payload.get("diagnosticos", [])
    chart_b64 = payload.get("chart_png_base64", "")
    area_lower = area.lower()

    wb = Workbook()

    # Styles
    header_fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    title_font = Font(bold=True, size=14, color="2C3E50")
    subtitle_font = Font(italic=True, size=10, color="666666")
    thin = Side(border_style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")

    # ═══ Sheet 1: RESUMO ═══
    ws1 = wb.active
    ws1.title = "Resumo"
    ws1["A1"] = "XTRI ENGINE — Relatorio de Correcao TRI"
    ws1["A1"].font = title_font
    ws1.merge_cells("A1:D1")
    ws1["A2"] = f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws1["A2"].font = subtitle_font
    ws1.merge_cells("A2:D2")

    ws1["A4"] = "Parametros da Correcao"
    ws1["A4"].font = Font(bold=True, size=12, color="2C3E50")
    params = [
        ("Area", area),
        ("Ano de Referencia ENEM", ano),
        ("Metodo de Ancoragem", metodo),
        ("Total de Alunos", n_alunos),
    ]
    for i, (k, v) in enumerate(params, start=5):
        ws1[f"A{i}"] = k
        ws1[f"A{i}"].font = Font(bold=True)
        ws1[f"B{i}"] = v

    ws1["A10"] = "Estatisticas da Turma"
    ws1["A10"].font = Font(bold=True, size=12, color="2C3E50")
    if diagnosticos:
        d = diagnosticos[0]
        stats = [
            ("Nota Media", d.get("nota_media", "")),
            ("Nota Minima", d.get("nota_min", "")),
            ("Nota Maxima", d.get("nota_max", "")),
            ("b medio (dificuldade media)", d.get("b_medio", "")),
            ("b desvio padrao", d.get("b_desvio", "")),
            ("% dentro do intervalo ENEM", f"{d.get('pct_dentro_intervalo', 0)}%"),
            ("Itens faceis (b < -1.5)", d.get("itens_faceis", 0)),
            ("Itens medios", d.get("itens_medios", 0)),
            ("Itens dificeis (b > 1.5)", d.get("itens_dificeis", 0)),
            ("acert_dif medio (Primi)", d.get("acert_dif_medio", 0)),
        ]
        for i, (k, v) in enumerate(stats, start=11):
            ws1[f"A{i}"] = k
            ws1[f"A{i}"].font = Font(bold=True)
            ws1[f"B{i}"] = v

    niveis_count: Dict[str, int] = {}
    for row in resultados:
        estado = row.get(f"estado_{area_lower}")
        if estado:
            niveis_count[estado] = niveis_count.get(estado, 0) + 1

    ws1["A23"] = "Distribuicao por Nivel Pedagogico"
    ws1["A23"].font = Font(bold=True, size=12, color="2C3E50")
    ordem = ["AVANCADO", "PROFICIENTE", "BASICO", "ELEMENTAR", "INICIANTE"]
    for i, nivel in enumerate(ordem, start=24):
        ws1[f"A{i}"] = nivel
        ws1[f"A{i}"].font = Font(bold=True)
        ws1[f"B{i}"] = niveis_count.get(nivel, 0)
        if n_alunos > 0:
            ws1[f"C{i}"] = f"{niveis_count.get(nivel, 0) / n_alunos * 100:.1f}%"

    ws1.column_dimensions["A"].width = 35
    ws1.column_dimensions["B"].width = 18
    ws1.column_dimensions["C"].width = 12

    # ═══ Sheet 2: ALUNOS ═══
    ws2 = wb.create_sheet("Alunos")
    headers = ["ID", "Acertos", "Theta", "Nota TRI", "% Dificeis (acert_dif)", "Nivel"]
    for col_idx, header in enumerate(headers, start=1):
        cell = ws2.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = border

    for row_idx, row in enumerate(resultados, start=2):
        ws2.cell(row=row_idx, column=1, value=row.get("aluno_id", ""))
        ws2.cell(row=row_idx, column=2, value=row.get(f"acertos_{area_lower}"))
        ws2.cell(row=row_idx, column=3, value=row.get(f"theta_{area_lower}"))
        ws2.cell(row=row_idx, column=4, value=row.get(f"nota_{area_lower}"))
        ws2.cell(row=row_idx, column=5, value=row.get(f"acert_dif_{area_lower}"))
        ws2.cell(row=row_idx, column=6, value=row.get(f"estado_{area_lower}"))
        for c in range(1, 7):
            ws2.cell(row=row_idx, column=c).border = border
            ws2.cell(row=row_idx, column=c).alignment = center if c > 1 else left

    ws2.column_dimensions["A"].width = 28
    ws2.column_dimensions["B"].width = 10
    ws2.column_dimensions["C"].width = 12
    ws2.column_dimensions["D"].width = 12
    ws2.column_dimensions["E"].width = 22
    ws2.column_dimensions["F"].width = 15
    ws2.freeze_panes = "A2"

    # ═══ Sheet 3: TABELA DIFICULDADES ═══
    ws3 = wb.create_sheet("Tabela Dificuldades")
    ws3["A1"] = f"Tabela de Dificuldades — {area} / ENEM {ano}"
    ws3["A1"].font = title_font
    ws3.merge_cells("A1:G1")
    ws3["A2"] = "Comparacao entre a distribuicao do simulado e os limites oficiais ENEM"
    ws3["A2"].font = subtitle_font
    ws3.merge_cells("A2:G2")

    from collections import defaultdict

    freq_ac: Dict[int, int] = defaultdict(int)
    notas_ac: Dict[int, list] = defaultdict(list)
    for row in resultados:
        acertos = row.get(f"acertos_{area_lower}")
        nota = row.get(f"nota_{area_lower}")
        if acertos is not None and nota is not None:
            freq_ac[acertos] += 1
            notas_ac[acertos].append(nota)

    table_headers = [
        "Acertos",
        "Freq Simulado",
        "Ref Min",
        "Ref Med",
        "Ref Max",
        "Nota Sim (media)",
        "Nota Sim (min-max)",
    ]
    for col_idx, header in enumerate(table_headers, start=1):
        cell = ws3.cell(row=4, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = border

    for i, ponto in enumerate(referencia, start=5):
        acertos = ponto["acertos"]
        notas = notas_ac.get(acertos, [])
        ws3.cell(row=i, column=1, value=acertos)
        ws3.cell(row=i, column=2, value=freq_ac.get(acertos, 0))
        ws3.cell(row=i, column=3, value=ponto["tri_min"])
        ws3.cell(row=i, column=4, value=ponto["tri_med"])
        ws3.cell(row=i, column=5, value=ponto["tri_max"])
        if notas:
            ws3.cell(row=i, column=6, value=round(sum(notas) / len(notas), 1))
            ws3.cell(row=i, column=7, value=f"{min(notas):.0f} - {max(notas):.0f}")
        for c in range(1, 8):
            ws3.cell(row=i, column=c).border = border
            ws3.cell(row=i, column=c).alignment = center

    ws3.column_dimensions["A"].width = 10
    ws3.column_dimensions["B"].width = 14
    ws3.column_dimensions["C"].width = 12
    ws3.column_dimensions["D"].width = 12
    ws3.column_dimensions["E"].width = 12
    ws3.column_dimensions["F"].width = 18
    ws3.column_dimensions["G"].width = 20
    ws3.freeze_panes = "A5"

    # ═══ Sheet 4: GRAFICO (if chart image provided) ═══
    if chart_b64:
        try:
            from openpyxl.drawing.image import Image as XLImage

            if "," in chart_b64:
                chart_b64 = chart_b64.split(",", 1)[1]
            img_bytes = base64.b64decode(chart_b64)
            img_stream = io.BytesIO(img_bytes)
            ws4 = wb.create_sheet("Grafico")
            ws4["A1"] = f"Grafico de Ancoragem — {area} / ENEM {ano}"
            ws4["A1"].font = title_font
            ws4.merge_cells("A1:H1")
            xlimg = XLImage(img_stream)
            xlimg.width = 900
            xlimg.height = 520
            ws4.add_image(xlimg, "A3")
        except Exception as exc:  # noqa: BLE001
            logger.warning("excel: failed to embed chart image: %s", exc)

    buf = io.BytesIO()
    wb.save(buf)
    filename = f"xtri_{area}_{ano}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return buf.getvalue(), filename


def _build_validator_csv_bytes(payload: Dict[str, Any]) -> tuple[bytes, str]:
    """Gera o CSV da comparação V3 x V4 do validator."""

    from datetime import datetime

    area = str(payload.get("area", "")).upper()
    ano = int(payload.get("ano", 0) or 0)
    resultados = payload.get("resultados", [])
    if not area or not isinstance(resultados, list):
        raise ValueError("Payload inválido para exportação do validator")

    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";")
    writer.writerow(
        [
            "ID",
            "Acertos",
            "Theta_V3",
            "Theta_V4",
            "Nota_V3",
            "Nota_V4",
            "Diff",
            "Estado_V3",
            "Estado_V4",
            "Match",
        ]
    )
    for row in resultados:
        writer.writerow(
            [
                row.get("aluno_id", ""),
                row.get("acertos", ""),
                row.get("theta_v3", ""),
                row.get("theta_v4", ""),
                row.get("nota_v3", ""),
                row.get("nota_v4", ""),
                row.get("diff", ""),
                row.get("estado_v3", ""),
                row.get("estado_v4", ""),
                "SIM" if row.get("estado_v3") == row.get("estado_v4") else "NAO",
            ]
        )

    csv_bytes = ("\ufeff" + buffer.getvalue()).encode("utf-8")
    filename = (
        f"xtri_validator_{area}_{ano}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    )
    return csv_bytes, filename


def _build_validator_excel_bytes(payload: Dict[str, Any]) -> tuple[bytes, str]:
    """Gera o XLSX da comparação V3 x V4 para uso no validator."""

    from datetime import datetime

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    area = str(payload.get("area", "")).upper()
    ano = int(payload.get("ano", 0) or 0)
    co_prova = str(payload.get("co_prova", ""))
    estatisticas = payload.get("estatisticas", {}) or {}
    transformacao = payload.get("transformacao_inep", {}) or {}
    resultados = payload.get("resultados", []) or []
    avisos = payload.get("avisos", []) or []

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Resumo"

    header_fill = PatternFill(start_color="1E3C72", end_color="1E3C72", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    title_font = Font(bold=True, size=14, color="1E3C72")
    subtitle_font = Font(italic=True, size=10, color="666666")
    thin = Side(border_style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")

    ws1["A1"] = "XTRI Validator — Comparação V3 vs V4"
    ws1["A1"].font = title_font
    ws1.merge_cells("A1:D1")
    ws1["A2"] = f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws1["A2"].font = subtitle_font
    ws1.merge_cells("A2:D2")

    params = [
        ("Ano ENEM", ano),
        ("Área", area),
        ("CO_PROVA", co_prova),
        ("Alunos", estatisticas.get("n_alunos", len(resultados))),
        ("Itens", payload.get("n_itens", "")),
        ("Tempo (s)", payload.get("tempo_processamento", "")),
    ]
    ws1["A4"] = "Parâmetros"
    ws1["A4"].font = Font(bold=True, size=12, color="1E3C72")
    for i, (k, v) in enumerate(params, start=5):
        ws1[f"A{i}"] = k
        ws1[f"A{i}"].font = Font(bold=True)
        ws1[f"B{i}"] = v

    metrics = [
        ("Correlação", estatisticas.get("correlacao", "")),
        ("MAE", estatisticas.get("mae", "")),
        ("RMSE", estatisticas.get("rmse", "")),
        ("Max |diff|", estatisticas.get("max_diff_abs", "")),
        ("Concordância de nível (%)", estatisticas.get("pct_concordancia_nivel", "")),
        ("Média V3", estatisticas.get("nota_media_v3", "")),
        ("Média V4", estatisticas.get("nota_media_v4", "")),
        ("Sigma INEP", transformacao.get("sigma", "")),
        ("Mu INEP", transformacao.get("mu", "")),
    ]
    ws1["A13"] = "Métricas"
    ws1["A13"].font = Font(bold=True, size=12, color="1E3C72")
    for i, (k, v) in enumerate(metrics, start=14):
        ws1[f"A{i}"] = k
        ws1[f"A{i}"].font = Font(bold=True)
        ws1[f"B{i}"] = v

    if avisos:
        ws1["A25"] = "Avisos"
        ws1["A25"].font = Font(bold=True, size=12, color="1E3C72")
        for i, warning in enumerate(avisos, start=26):
            ws1[f"A{i}"] = f"• {warning}"
            ws1.merge_cells(start_row=i, start_column=1, end_row=i, end_column=4)

    ws1.column_dimensions["A"].width = 34
    ws1.column_dimensions["B"].width = 18
    ws1.column_dimensions["C"].width = 12
    ws1.column_dimensions["D"].width = 12

    ws2 = wb.create_sheet("Comparacao")
    headers = [
        "ID",
        "Acertos",
        "Theta V3",
        "Theta V4",
        "Nota V3",
        "Nota V4",
        "Diff",
        "Nivel V3",
        "Nivel V4",
        "Match",
    ]
    for col_idx, header in enumerate(headers, start=1):
        cell = ws2.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = border

    for row_idx, row in enumerate(resultados, start=2):
        values = [
            row.get("aluno_id", ""),
            row.get("acertos", ""),
            row.get("theta_v3", ""),
            row.get("theta_v4", ""),
            row.get("nota_v3", ""),
            row.get("nota_v4", ""),
            row.get("diff", ""),
            row.get("estado_v3", ""),
            row.get("estado_v4", ""),
            "SIM" if row.get("estado_v3") == row.get("estado_v4") else "NAO",
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws2.cell(row=row_idx, column=col_idx, value=value)
            cell.border = border
            cell.alignment = center

    for col in ("A", "B", "C", "D", "E", "F", "G", "H", "I", "J"):
        ws2.column_dimensions[col].width = 16 if col != "A" else 28
    ws2.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    filename = (
        f"xtri_validator_{area}_{ano}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    )
    return buf.getvalue(), filename


@router.post("/api/v3/exportar-csv")
async def exportar_csv(payload: Dict[str, Any] = Body(...), salvar_localmente: bool = False):
    """Exporta o CSV do motor principal; no app local pode salvar em Downloads."""

    try:
        csv_bytes, filename = _build_resultados_csv_bytes(payload)
        if salvar_localmente:
            return JSONResponse(_save_local_export(filename, csv_bytes))
        return StreamingResponse(
            io.BytesIO(csv_bytes),
            media_type="text/csv; charset=utf-8",
            headers={
                "Content-Disposition": f'attachment; filename="{filename}"',
                "Content-Length": str(len(csv_bytes)),
            },
        )
    except Exception as exc:  # noqa: BLE001
        logger.exception("exportar-csv failed")
        raise HTTPException(status_code=500, detail=f"Erro ao gerar CSV: {exc}")


@router.post("/api/v3/exportar-excel")
async def exportar_excel(payload: Dict[str, Any] = Body(...), salvar_localmente: bool = False):
    """
    Generate an XLSX file with all results: students, difficulty table,
    diagnostics, and the anchor chart as an embedded image.

    When ``salvar_localmente`` is true, writes the file to Downloads and
    returns the destination path instead of streaming it through the WebView.
    """
    try:
        xlsx_bytes, filename = _build_resultados_excel_bytes(payload)
        if salvar_localmente:
            return JSONResponse(_save_local_export(filename, xlsx_bytes))
        return StreamingResponse(
            io.BytesIO(xlsx_bytes),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": f'attachment; filename="{filename}"',
                "Content-Length": str(len(xlsx_bytes)),
            },
        )
    except Exception as exc:  # noqa: BLE001
        logger.exception("exportar-excel failed")
        raise HTTPException(status_code=500, detail=f"Erro ao gerar Excel: {exc}")


@router.post("/api/v4/enem/exportar-validacao-csv")
async def exportar_validacao_csv(
    payload: Dict[str, Any] = Body(...), salvar_localmente: bool = False
):
    """Exporta o comparativo do validator em CSV."""

    try:
        csv_bytes, filename = _build_validator_csv_bytes(payload)
        if salvar_localmente:
            return JSONResponse(_save_local_export(filename, csv_bytes))
        return StreamingResponse(
            io.BytesIO(csv_bytes),
            media_type="text/csv; charset=utf-8",
            headers={
                "Content-Disposition": f'attachment; filename="{filename}"',
                "Content-Length": str(len(csv_bytes)),
            },
        )
    except Exception as exc:  # noqa: BLE001
        logger.exception("exportar-validacao-csv failed")
        raise HTTPException(status_code=500, detail=f"Erro ao gerar CSV do validator: {exc}")


@router.post("/api/v4/enem/exportar-validacao-excel")
async def exportar_validacao_excel(
    payload: Dict[str, Any] = Body(...), salvar_localmente: bool = False
):
    """Exporta o comparativo do validator em XLSX."""

    try:
        xlsx_bytes, filename = _build_validator_excel_bytes(payload)
        if salvar_localmente:
            return JSONResponse(_save_local_export(filename, xlsx_bytes))
        return StreamingResponse(
            io.BytesIO(xlsx_bytes),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": f'attachment; filename="{filename}"',
                "Content-Length": str(len(xlsx_bytes)),
            },
        )
    except Exception as exc:  # noqa: BLE001
        logger.exception("exportar-validacao-excel failed")
        raise HTTPException(
            status_code=500, detail=f"Erro ao gerar Excel do validator: {exc}"
        )
