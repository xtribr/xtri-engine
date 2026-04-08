"""Tests for :mod:`app.microdata.excel_export` and the v5 export endpoint."""

from __future__ import annotations

import io

import numpy as np
import pandas as pd
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.microdata.cache import cache_key
from app.microdata.calibration import calibration_summary
from app.microdata.excel_export import build_microdata_csv, build_microdata_excel
from app.microdata.loader import default_csv_path
from app.microdata.schema import AREAS

XLSX_MAGIC = b"PK\x03\x04"


@pytest.fixture
def calibrated_synthetic_df():
    """Tiny but realistic calibrated DataFrame for unit-testing the writer."""
    rng = np.random.default_rng(99)
    n = 200
    data: dict = {"NU_SEQUENCIAL": [str(1_000_000 + i) for i in range(n)]}
    data["SG_UF_ESC"] = rng.choice(["SP", "RJ", "MG", "BA", "CE"], size=n)
    data["NO_MUNICIPIO_ESC"] = rng.choice(
        ["SAO PAULO", "RIO DE JANEIRO", "BELO HORIZONTE", "SALVADOR", "FORTALEZA"], size=n
    )
    data["CO_ESCOLA"] = rng.integers(35_000_000, 35_999_999, size=n)
    for area in AREAS:
        al = area.lower()
        ac = rng.integers(0, 46, size=n)
        ref = 200 + 13 * ac + rng.normal(0, 25, size=n)
        ref = np.clip(ref, 200, 1000)
        v3 = ref + rng.normal(-5, 30, size=n)
        v3 = np.clip(v3, 200, 1000)
        data[f"acertos_v3_{al}"] = ac
        data[f"theta_v3_{al}"] = (v3 - v3.mean()) / v3.std()
        data[f"se_v3_{al}"] = np.full(n, 0.3)
        data[f"nota_v3_{al}"] = v3
        data[f"nota_ref_{al}"] = ref
        data[f"estado_v3_{al}"] = np.where(v3 >= 600, "PROFICIENTE", "BASICO")
        data[f"acert_dif_{al}"] = rng.uniform(0, 1, size=n)
    return pd.DataFrame(data)


# ─────────────────────────────────────────────────────────────────────────────
# Excel writer unit tests
# ─────────────────────────────────────────────────────────────────────────────


class TestBuildExcel:
    @pytest.mark.unit
    def test_returns_valid_xlsx_bytes(self, calibrated_synthetic_df):
        summaries = calibration_summary(calibrated_synthetic_df)
        xlsx = build_microdata_excel(calibrated_synthetic_df, year=2024, summaries=summaries)
        assert isinstance(xlsx, bytes)
        assert xlsx.startswith(XLSX_MAGIC)
        assert len(xlsx) > 10_000

    @pytest.mark.unit
    def test_workbook_has_expected_sheets(self, calibrated_synthetic_df):
        summaries = calibration_summary(calibrated_synthetic_df)
        xlsx = build_microdata_excel(calibrated_synthetic_df, year=2024, summaries=summaries)
        wb = load_workbook(io.BytesIO(xlsx))
        names = wb.sheetnames
        assert "Resumo" in names
        assert "Relatorio Completo" in names
        # Default excludes Alunos sheet (lightweight mode)
        assert "Alunos" not in names
        # Per-area sheets are named after the Portuguese label
        assert any("Linguagens" in n for n in names)
        assert any("Ciências Humanas" in n for n in names)
        assert any("Ciências da Natureza" in n for n in names)
        assert any("Matemática" in n for n in names)

    @pytest.mark.unit
    def test_default_excludes_alunos_sheet(self, calibrated_synthetic_df):
        """Default Excel build is lightweight — no per-student sheet."""
        summaries = calibration_summary(calibrated_synthetic_df)
        xlsx = build_microdata_excel(calibrated_synthetic_df, year=2024, summaries=summaries)
        wb = load_workbook(io.BytesIO(xlsx))
        assert "Alunos" not in wb.sheetnames

    @pytest.mark.unit
    def test_include_students_adds_alunos_sheet(self, calibrated_synthetic_df):
        """Opt-in flag adds the per-student INEP-only sheet."""
        summaries = calibration_summary(calibrated_synthetic_df)
        xlsx = build_microdata_excel(
            calibrated_synthetic_df,
            year=2024,
            summaries=summaries,
            include_students=True,
        )
        wb = load_workbook(io.BytesIO(xlsx))
        assert "Alunos" in wb.sheetnames
        ws = wb["Alunos"]
        headers = [c.value for c in ws[1]]
        assert "NU_SEQUENCIAL" in headers
        assert "SG_UF_ESC" in headers
        for area in AREAS:
            al = area.lower()
            assert f"acertos_{al}" in headers
            assert f"nota_ref_{al}" in headers
            assert f"acert_dif_{al}" in headers
            # V3 columns must be ABSENT (microdata page is INEP-only)
            assert f"nota_v3_{al}" not in headers
            assert f"diff_{al}" not in headers
            assert f"theta_v3_{al}" not in headers
            assert f"estado_v3_{al}" not in headers
        # Row count = 1 header + n students
        assert ws.max_row == len(calibrated_synthetic_df) + 1

    @pytest.mark.unit
    def test_resumo_sheet_has_inep_descriptive_stats(self, calibrated_synthetic_df):
        summaries = calibration_summary(calibrated_synthetic_df)
        xlsx = build_microdata_excel(calibrated_synthetic_df, year=2024, summaries=summaries)
        wb = load_workbook(io.BytesIO(xlsx))
        ws = wb["Resumo"]
        # Title in row 1
        assert "XTRI" in str(ws["A1"].value)
        # Headers in row 4, area data starting row 5
        assert ws["A4"].value == "Área"
        assert ws["D4"].value == "Mean"
        assert ws["E4"].value == "Median"
        assert ws["A5"].value == "LC"
        assert ws["A8"].value == "MT"
        # Mean / Median columns populated
        assert ws["D5"].value is not None
        assert isinstance(ws["D5"].value, (int, float))
        assert ws["E5"].value is not None
        assert isinstance(ws["E5"].value, (int, float))


# ─────────────────────────────────────────────────────────────────────────────
# Endpoint tests
# ─────────────────────────────────────────────────────────────────────────────


class TestExcelEndpoint:
    @pytest.fixture
    def client(self):
        return TestClient(app)

    @pytest.mark.api
    def test_export_excel_for_missing_csv_returns_404(self, client):
        r = client.get("/api/v5/microdata/export-excel/9999")
        assert r.status_code == 404

    @pytest.mark.api
    def test_export_excel_default_is_lightweight(self, client):
        """Real-data smoke: default Excel skips Alunos sheet → fast + small."""
        csv_path = default_csv_path(2024)
        if not csv_path.is_file():
            pytest.skip("MICRODADOS_ENEM_2024.csv not on disk")
        raw_key = cache_key(csv_path, year=2024, sample_size=200_000, rng_seed=42)
        if not raw_key.with_stage("calibrated").parquet_path.is_file():
            pytest.skip("calibrated cache not present")

        r = client.get("/api/v5/microdata/export-excel/2024")
        assert r.status_code == 200
        assert "xtri_microdados_2024.xlsx" in r.headers["content-disposition"]
        body = r.content
        assert body.startswith(XLSX_MAGIC)
        # Default mode is ~860 KB (no Alunos sheet)
        assert 100_000 < len(body) < 5_000_000
        wb = load_workbook(io.BytesIO(body))
        assert "Resumo" in wb.sheetnames
        assert "Alunos" not in wb.sheetnames

    @pytest.mark.api
    def test_export_excel_with_students_includes_alunos(self, client):
        """Opt-in include_students=true adds the heavy per-student sheet."""
        csv_path = default_csv_path(2024)
        if not csv_path.is_file():
            pytest.skip("MICRODADOS_ENEM_2024.csv not on disk")
        raw_key = cache_key(csv_path, year=2024, sample_size=200_000, rng_seed=42)
        if not raw_key.with_stage("calibrated").parquet_path.is_file():
            pytest.skip("calibrated cache not present")

        r = client.get("/api/v5/microdata/export-excel/2024?include_students=true")
        assert r.status_code == 200
        assert "xtri_microdados_2024_full.xlsx" in r.headers["content-disposition"]
        body = r.content
        assert body.startswith(XLSX_MAGIC)
        assert len(body) > 5_000_000  # full version >5 MB
        wb = load_workbook(io.BytesIO(body))
        assert "Alunos" in wb.sheetnames
        assert wb["Alunos"].max_row > 100_000

    @pytest.mark.api
    def test_export_csv_default(self, client):
        csv_path = default_csv_path(2024)
        if not csv_path.is_file():
            pytest.skip("MICRODADOS_ENEM_2024.csv not on disk")
        raw_key = cache_key(csv_path, year=2024, sample_size=200_000, rng_seed=42)
        if not raw_key.with_stage("calibrated").parquet_path.is_file():
            pytest.skip("calibrated cache not present")

        r = client.get("/api/v5/microdata/export-csv/2024")
        assert r.status_code == 200
        assert r.headers["content-type"].startswith("text/csv")
        assert "xtri_microdados_alunos_2024.csv" in r.headers["content-disposition"]
        body = r.content
        # UTF-8 BOM
        assert body.startswith(b"\xef\xbb\xbf")
        # Header row contains expected columns — slice up to first newline
        first_line_end = body.find(b"\n", 3)
        assert first_line_end > 0
        first_line = body[3:first_line_end].decode("utf-8", errors="replace")
        assert "NU_SEQUENCIAL" in first_line
        # INEP-only columns
        assert "nota_ref_mt" in first_line
        assert "acert_dif_mt" in first_line
        # V3 columns must NOT be present
        assert "nota_v3_mt" not in first_line
        assert "diff_mt" not in first_line
        # Many rows (~200K)
        assert body.count(b"\n") > 100_000

    @pytest.mark.api
    def test_export_csv_for_missing_csv_returns_404(self, client):
        r = client.get("/api/v5/microdata/export-csv/9999")
        assert r.status_code == 404
