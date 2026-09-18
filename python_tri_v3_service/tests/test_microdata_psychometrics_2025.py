"""Testes do relatório psicométrico complementar do ENEM 2025."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from app.microdata.psychometrics_2025 import (
    analyze_enem_2025_area,
    export_psychometric_report,
    resolve_enem_2025_paths,
)


ITEM_COLUMNS = [
    "CO_POSICAO",
    "SG_AREA",
    "CO_ITEM",
    "TX_GABARITO",
    "CO_HABILIDADE",
    "IN_ITEM_ABAN",
    "TX_MOTIVO_ABAN",
    "NU_PARAM_A",
    "NU_PARAM_B",
    "NU_PARAM_C",
    "TX_COR",
    "CO_PROVA",
    "TP_LINGUA",
    "IN_ITEM_ADAPTADO",
]


def _write_csv(frame: pd.DataFrame, path: Path) -> Path:
    frame.to_csv(path, sep=";", encoding="latin-1", index=False)
    return path


def _mt_items() -> pd.DataFrame:
    rows = []
    keys = "ABCDE" * 9
    for offset, key in enumerate(keys):
        rows.append(
            {
                "CO_POSICAO": 136 + offset,
                "SG_AREA": "MT",
                "CO_ITEM": str(10_000 + offset),
                "TX_GABARITO": key,
                "CO_HABILIDADE": str(1 + offset % 5),
                "IN_ITEM_ABAN": "0",
                "TX_MOTIVO_ABAN": "",
                "NU_PARAM_A": "1,20",
                "NU_PARAM_B": "0,30",
                "NU_PARAM_C": "0,20",
                "TX_COR": "AZUL",
                "CO_PROVA": "1471",
                "TP_LINGUA": "",
                "IN_ITEM_ADAPTADO": "0",
            }
        )
    return pd.DataFrame(rows, columns=ITEM_COLUMNS)


def _wrong(key: str) -> str:
    return "A" if key != "A" else "B"


def _mt_results() -> pd.DataFrame:
    key = "ABCDE" * 9
    nested_30 = key[:30] + "".join(_wrong(char) for char in key[30:])
    nested_15 = key[:15] + "".join(_wrong(char) for char in key[15:])
    return pd.DataFrame(
        [
            {"TP_PRESENCA_MT": "1", "CO_PROVA_MT": "1471", "TX_RESPOSTAS_MT": key},
            {"TP_PRESENCA_MT": "1", "CO_PROVA_MT": "1471", "TX_RESPOSTAS_MT": nested_30},
            {"TP_PRESENCA_MT": "1", "CO_PROVA_MT": "1471", "TX_RESPOSTAS_MT": nested_15},
            {"TP_PRESENCA_MT": "0", "CO_PROVA_MT": "1471", "TX_RESPOSTAS_MT": key},
            {"TP_PRESENCA_MT": "1", "CO_PROVA_MT": "1607", "TX_RESPOSTAS_MT": key},
            {"TP_PRESENCA_MT": "1", "CO_PROVA_MT": "1471", "TX_RESPOSTAS_MT": key[:-1]},
        ]
    )


@pytest.mark.unit
def test_resolve_split_2025_package(tmp_path):
    data = tmp_path / "DADOS"
    data.mkdir()
    (data / "RESULTADOS_2025.csv").write_text("x\n", encoding="latin-1")
    (data / "ITENS_PROVA_2025.csv").write_text("x\n", encoding="latin-1")
    (data / "PARTICIPANTES_2025.csv").write_text("x\n", encoding="latin-1")

    paths = resolve_enem_2025_paths(tmp_path)

    assert paths.root == tmp_path.resolve()
    assert paths.resultados == (data / "RESULTADOS_2025.csv").resolve()
    assert paths.itens == (data / "ITENS_PROVA_2025.csv").resolve()
    assert paths.participantes == (data / "PARTICIPANTES_2025.csv").resolve()


@pytest.mark.unit
def test_mt_report_filters_presence_bam_and_invalid_length(tmp_path):
    items_path = _write_csv(_mt_items(), tmp_path / "ITENS_PROVA_2025.csv")
    results_path = _write_csv(_mt_results(), tmp_path / "RESULTADOS_2025.csv")

    report = analyze_enem_2025_area(
        results_path=results_path,
        items_path=items_path,
        area="MT",
        application="regular",
        chunksize=2,
    )

    assert report.summary["linhas_lidas"] == 6
    assert report.summary["participantes_presentes_area"] == 5
    assert report.summary["participantes_analisados"] == 3
    assert report.summary["participantes_por_aplicacao"] == {"bam": 1, "regular": 4}
    assert report.summary["exclusoes"] == {"resposta_tamanho_invalido": 1}
    assert report.summary["itens_canonicos"] == 45
    assert report.summary["itens_analisados"] == 45
    assert report.forms.loc[0, "n_analisados"] == 3
    assert report.forms.loc[0, "media_acertos"] == pytest.approx(30.0)
    assert 0 <= report.forms.loc[0, "alpha_cronbach"] <= 1

    item = report.items.loc[report.items["co_item"].eq("10020")].iloc[0]
    assert item["taxa_acerto"] == pytest.approx(2 / 3)
    assert item["correlacao_item_total_corrigida"] > 0
    counts = report.distractors.loc[report.distractors["co_item"].eq("10020"), "n"].sum()
    assert counts == 3


@pytest.mark.unit
def test_lc_uses_language_specific_first_five_items(tmp_path):
    rows = []
    for language, keys, base in [(0, "ABCDE", 20_000), (1, "EDCBA", 30_000)]:
        for position, key in enumerate(keys, start=1):
            rows.append(
                {
                    "CO_POSICAO": position,
                    "SG_AREA": "LC",
                    "CO_ITEM": str(base + position),
                    "TX_GABARITO": key,
                    "CO_HABILIDADE": "1",
                    "IN_ITEM_ABAN": "0",
                    "TX_MOTIVO_ABAN": "",
                    "NU_PARAM_A": "1,20",
                    "NU_PARAM_B": "0,30",
                    "NU_PARAM_C": "0,20",
                    "TX_COR": "AZUL",
                    "CO_PROVA": "1459",
                    "TP_LINGUA": str(language),
                    "IN_ITEM_ADAPTADO": "0",
                }
            )
    common_keys = "ABCDE" * 8
    for offset, key in enumerate(common_keys, start=6):
        rows.append(
            {
                "CO_POSICAO": offset,
                "SG_AREA": "LC",
                "CO_ITEM": str(40_000 + offset),
                "TX_GABARITO": key,
                "CO_HABILIDADE": "2",
                "IN_ITEM_ABAN": "0",
                "TX_MOTIVO_ABAN": "",
                "NU_PARAM_A": "1,30",
                "NU_PARAM_B": "0,20",
                "NU_PARAM_C": "0,20",
                "TX_COR": "AZUL",
                "CO_PROVA": "1459",
                "TP_LINGUA": "",
                "IN_ITEM_ADAPTADO": "0",
            }
        )
    items_path = _write_csv(
        pd.DataFrame(rows, columns=ITEM_COLUMNS), tmp_path / "ITENS_PROVA_2025.csv"
    )
    results_path = _write_csv(
        pd.DataFrame(
            [
                {"TP_PRESENCA_LC": "1", "CO_PROVA_LC": "1459", "TX_RESPOSTAS_LC": "ABCDE" + common_keys, "TP_LINGUA": "0"},
                {"TP_PRESENCA_LC": "1", "CO_PROVA_LC": "1459", "TX_RESPOSTAS_LC": "EDCBA" + common_keys, "TP_LINGUA": "1"},
            ]
        ),
        tmp_path / "RESULTADOS_2025.csv",
    )

    report = analyze_enem_2025_area(
        results_path=results_path,
        items_path=items_path,
        area="LC",
        chunksize=1,
    )

    assert report.summary["participantes_analisados"] == 2
    assert report.summary["itens_canonicos"] == 50
    english = report.items.loc[report.items["co_item"].eq("20001")].iloc[0]
    spanish = report.items.loc[report.items["co_item"].eq("30001")].iloc[0]
    common = report.items.loc[report.items["co_item"].eq("40006")].iloc[0]
    assert english["n_expostos"] == 1
    assert spanish["n_expostos"] == 1
    assert common["n_expostos"] == 2


@pytest.mark.unit
def test_export_contains_auditable_sheets(tmp_path):
    items_path = _write_csv(_mt_items(), tmp_path / "ITENS_PROVA_2025.csv")
    results_path = _write_csv(_mt_results(), tmp_path / "RESULTADOS_2025.csv")
    report = analyze_enem_2025_area(
        results_path=results_path,
        items_path=items_path,
        area="MT",
    )
    output = tmp_path / "relatorio.xlsx"
    summary = tmp_path / "resumo.json"

    export_psychometric_report(report, output, output_json=summary)

    workbook = load_workbook(output, read_only=True, data_only=True)
    assert workbook.sheetnames == ["Resumo", "Itens", "Distratores", "Cadernos", "Qualidade"]
    assert workbook["Itens"].max_row == 46
    assert summary.is_file()
