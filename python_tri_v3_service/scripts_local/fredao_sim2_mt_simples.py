"""Fredão Sim2 — relatório SIMPLES de Matemática.

Excel com 4 colunas apenas:
  1. Nome (ID do aluno)
  2. Acertos (sobre os 44 itens válidos — Q148 excluída)
  3. TRI 2017 — V4 híbrido (a,b,c oficial INEP 2017 + tabela pedagógica 2017)
  4. TRI 2024 — V4 puro (a,b,c oficial INEP 2024 + σ·θ+μ INEP 2024)

Ordenado por TRI 2024 desc (ranking).
"""

from __future__ import annotations

import json
import sys
import time
from pathlib import Path
from typing import Dict, List, Tuple

import numpy as np
import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from app.irt.scoring import eap_theta  # noqa: E402
from app.irt.transform import ancorar_notas  # noqa: E402
from app.tabela_referencia import TabelaReferenciaTRI  # noqa: E402

MT_PATH = Path("/Users/home/Library/CloudStorage/Dropbox/Fredao Simulados/Matemática - Simulado 2.xlsx")
PARAMS_JSON = Path(__file__).parent / "data" / "params_oficiais.json"
MT_ANULADA_IDX = 12  # Q148

SIGMA_MT_2024, MU_MT_2024 = 129.6484, 500.0186

OUT_DIR = Path("/Users/home/Downloads/XTRI Engine")
OUT_DIR.mkdir(parents=True, exist_ok=True)


def carregar_mt() -> Tuple[List[str], np.ndarray]:
    wb = openpyxl.load_workbook(MT_PATH, data_only=True)
    ws = wb.worksheets[0]
    ids, rows = [], []
    for r in range(2, ws.max_row + 1):
        aluno = ws.cell(r, 1).value
        if not aluno:
            continue
        ids.append(str(aluno))
        rows.append([int(ws.cell(r, c).value or 0) for c in range(2, 47)])
    return ids, np.array(rows, dtype=int)


def params_para_dicts(plist: List[List]) -> List[Dict]:
    return [{"a": float(p[1]), "b": float(p[2]), "c": float(p[3])} for p in plist]


def estimar_thetas(resp: np.ndarray, params: List[Dict]) -> Tuple[np.ndarray, np.ndarray]:
    n = len(resp)
    thetas = np.zeros(n)
    ses = np.zeros(n)
    for i, row in enumerate(resp):
        t, se = eap_theta(row, params)
        thetas[i] = t
        ses[i] = se
    return thetas, ses


def main():
    t0 = time.time()
    print("FREDÃO Sim2 — MT relatório simples (4 colunas)")
    print("=" * 60)

    ids, mat_full = carregar_mt()
    keep = [j for j in range(mat_full.shape[1]) if j != MT_ANULADA_IDX]
    mat = mat_full[:, keep]
    acertos = mat.sum(axis=1).astype(int)
    print(f"  Alunos: {len(ids)} | itens válidos: {mat.shape[1]} (Q148 excluída)")

    with open(PARAMS_JSON) as f:
        params_of = json.load(f)

    # Excluir Q148 dos params também (índice 12 nas listas originais)
    params_2017_full = params_para_dicts(params_of["2017"]["MT"])
    params_2024_full = params_para_dicts(params_of["2024"]["MT"])
    params_2017 = [params_2017_full[j] for j in keep]
    params_2024 = [params_2024_full[j] for j in keep]

    # === TRI 2017 — V4 híbrido (a,b,c oficial 2017 + tabela 2017) ===
    print("  Estimando θ com params INEP 2017...")
    thetas_2017, ses_2017 = estimar_thetas(mat, params_2017)
    tabelas = TabelaReferenciaTRI.carregar_multiano()
    thetas_list_2017 = [
        {
            "aluno_id": ids[i], "nome": ids[i], "acertos": int(acertos[i]),
            "theta": float(thetas_2017[i]), "se": float(ses_2017[i]),
            "acertos_faceis": 0, "acertos_medios": 0,
            "acertos_dificeis": 0, "acert_dif": 0.0,
        }
        for i in range(len(ids))
    ]
    res_2017 = ancorar_notas(thetas_list_2017, tabelas[2017], "MT")
    notas_2017 = np.array([r["nota_ancorada"] for r in res_2017])
    print(f"  TRI 2017: μ={notas_2017.mean():.1f} σ={notas_2017.std():.1f} "
          f"min={notas_2017.min():.0f} max={notas_2017.max():.0f}")

    # === TRI 2024 — V4 puro (a,b,c oficial 2024 + σ·θ+μ 2024) ===
    print("  Estimando θ com params INEP 2024...")
    thetas_2024, _ = estimar_thetas(mat, params_2024)
    notas_2024 = np.clip(SIGMA_MT_2024 * thetas_2024 + MU_MT_2024, 0.0, 1000.0)
    print(f"  TRI 2024: μ={notas_2024.mean():.1f} σ={notas_2024.std():.1f} "
          f"min={notas_2024.min():.0f} max={notas_2024.max():.0f}")

    # === DataFrame final (4 colunas) ===
    df = pd.DataFrame({
        "Nome": ids,
        "Acertos": acertos,
        "TRI 2017": np.round(notas_2017, 1),
        "TRI 2024": np.round(notas_2024, 1),
    })
    df = df.sort_values("TRI 2024", ascending=False).reset_index(drop=True)

    ts = time.strftime("%Y%m%d_%H%M%S")
    xlsx_path = OUT_DIR / f"xtri_FREDAO_sim2_MT_simples_{ts}.xlsx"

    df.to_excel(xlsx_path, sheet_name="MT", index=False)

    # Formatação leve
    wb = openpyxl.load_workbook(xlsx_path)
    from openpyxl.styles import Font, Alignment, PatternFill
    ws = wb["MT"]
    bold = Font(bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="2E86AB", end_color="2E86AB", fill_type="solid")
    centro = Alignment(horizontal="center", vertical="center")
    for cell in ws[1]:
        cell.font = bold
        cell.alignment = centro
        cell.fill = fill_header
    # Larguras
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 14
    # Centro nos dados
    for r in range(2, ws.max_row + 1):
        for c in range(1, 5):
            ws.cell(r, c).alignment = centro
    # Freeze top row
    ws.freeze_panes = "A2"
    wb.save(xlsx_path)

    print(f"\n  XLSX: {xlsx_path}")
    print(f"\nTOTAL: {time.time()-t0:.1f}s")

    # Print primeiros e últimos 5
    print("\nTop 5:")
    print(df.head(5).to_string(index=False))
    print("\nBottom 5:")
    print(df.tail(5).to_string(index=False))


if __name__ == "__main__":
    main()
