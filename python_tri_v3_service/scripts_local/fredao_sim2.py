"""Fredão Simulado 2 — Correção TRI V3 + escala σ·θ+μ INEP 2024.

- 465 alunos
- CN: 45 itens (Q91-135), 0 NaN, 0 anuladas
- MT: 45 itens (Q136-180), 0 NaN, Q148 ANULADA → excluída (44 itens válidos)
- σ/μ INEP 2024: CN 113.1042/501.1432, MT 129.6484/500.0186
"""

from __future__ import annotations

import os
import sys
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Tuple

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
import pandas as pd

# Permitir importar app.* sem instalação
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from app.irt.scoring import eap_theta  # noqa: E402
from app.irt.unidim_em import calibrar_itens_mml_em  # noqa: E402

DROPBOX = Path("/Users/home/Library/CloudStorage/Dropbox/Fredao Simulados")
CN_PATH = DROPBOX / "Natureza - Simulado 2.xlsx"
MT_PATH = DROPBOX / "Matemática - Simulado 2.xlsx"

# INEP 2024 oficial
SIGMA_CN, MU_CN = 113.1042, 501.1432
SIGMA_MT, MU_MT = 129.6484, 500.0186

# Q148 = índice 12 em MT (Q136 = idx 0). 0-indexed.
MT_ANULADA_IDX = 12  # corresponde a Q148

OUT_DIR = Path("/Users/home/Downloads/XTRI Engine")
OUT_DIR.mkdir(parents=True, exist_ok=True)


@dataclass(frozen=True)
class AreaResult:
    area: str
    n_alunos: int
    n_itens_validos: int
    ids: List[str]
    resp_matrix: np.ndarray
    item_params: List[Dict]
    thetas: np.ndarray
    ses: np.ndarray
    notas: np.ndarray
    acertos: np.ndarray


def carregar_planilha(path: Path, area: str) -> Tuple[List[str], np.ndarray, List[int]]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    header = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    # primeira col é ID; demais até 46 são as questões
    ids: List[str] = []
    rows: List[List[int]] = []
    q_numbers: List[int] = []
    for c, h in enumerate(header[1:46], 2):  # cols 2..46
        # header pode ser int (num da questão) ou string ("Anulada")
        if isinstance(h, str) and h.lower().startswith("anul"):
            q_numbers.append(-1)  # marca como anulada
        elif isinstance(h, (int, float)):
            q_numbers.append(int(h))
        else:
            q_numbers.append(-2)
    for r in range(2, ws.max_row + 1):
        aluno_id = ws.cell(r, 1).value
        if not aluno_id:
            continue
        ids.append(str(aluno_id))
        row = [int(ws.cell(r, c).value or 0) for c in range(2, 47)]
        rows.append(row)
    mat = np.array(rows, dtype=int)
    print(f"  [{area}] alunos={len(ids)}, itens={mat.shape[1]}, "
          f"acertos médios={mat.sum(axis=1).mean():.1f}/{mat.shape[1]}")
    return ids, mat, q_numbers


def calibrar_area(area: str, resp_matrix: np.ndarray) -> List[Dict]:
    print(f"  [{area}] calibrando MML/EM (c=0.20, n_quad=61)...")
    t0 = time.time()
    params = calibrar_itens_mml_em(resp_matrix, c=0.20, n_quad=61)
    elapsed = time.time() - t0
    a_vals = [p["a"] for p in params]
    b_vals = [p["b"] for p in params]
    saturados = sum(1 for p in params if p.get("saturado"))
    print(f"  [{area}] done in {elapsed:.1f}s. "
          f"a μ={np.mean(a_vals):.2f}±{np.std(a_vals):.2f}, "
          f"b μ={np.mean(b_vals):.2f}±{np.std(b_vals):.2f}, "
          f"saturados={saturados}")
    return params


def estimar_thetas(resp_matrix: np.ndarray, params: List[Dict]) -> Tuple[np.ndarray, np.ndarray]:
    thetas = np.zeros(len(resp_matrix))
    ses = np.zeros(len(resp_matrix))
    for i, row in enumerate(resp_matrix):
        t, se = eap_theta(row, params)
        thetas[i] = t
        ses[i] = se
    return thetas, ses


def aplicar_sigma_mu(thetas: np.ndarray, sigma: float, mu: float) -> np.ndarray:
    notas = sigma * thetas + mu
    return np.clip(notas, 0.0, 1000.0)


def nivel_pedagogico(nota: float) -> str:
    if nota >= 650:
        return "AVANÇADO"
    if nota >= 580:
        return "PROFICIENTE"
    if nota >= 480:
        return "BÁSICO"
    if nota >= 380:
        return "ELEMENTAR"
    return "INICIANTE"


def processar_area(area: str, path: Path, sigma: float, mu: float,
                   anulada_idx: int = -1) -> AreaResult:
    print(f"\n=== {area} ===")
    ids, mat_full, q_numbers = carregar_planilha(path, area)

    # Excluir item anulado se aplicável
    if anulada_idx >= 0:
        keep = [j for j in range(mat_full.shape[1]) if j != anulada_idx]
        mat = mat_full[:, keep]
        print(f"  [{area}] EXCLUINDO Q{q_numbers[anulada_idx] if q_numbers[anulada_idx] >0 else 148} "
              f"(anulada, idx {anulada_idx}) → matriz {mat.shape}")
    else:
        mat = mat_full

    params = calibrar_area(area, mat)
    thetas, ses = estimar_thetas(mat, params)
    notas = aplicar_sigma_mu(thetas, sigma, mu)
    acertos = mat.sum(axis=1)

    print(f"  [{area}] θ μ={thetas.mean():.3f}, σ={thetas.std():.3f}")
    print(f"  [{area}] Notas μ={notas.mean():.1f}, σ={notas.std():.1f}, "
          f"min={notas.min():.0f}, max={notas.max():.0f}")

    return AreaResult(area, len(ids), mat.shape[1], ids,
                      mat, params, thetas, ses, notas, acertos)


def flag_atipico(r_cn: AreaResult, r_mt: AreaResult, i: int) -> str:
    flags = []
    # |θ| muito alto → saturação informacional
    if abs(r_cn.thetas[i]) > 2.5:
        flags.append(f"θ_CN={r_cn.thetas[i]:+.2f}")
    if abs(r_mt.thetas[i]) > 2.5:
        flags.append(f"θ_MT={r_mt.thetas[i]:+.2f}")
    # SE muito alto → estimativa imprecisa
    if r_cn.ses[i] > 0.6:
        flags.append(f"SE_CN={r_cn.ses[i]:.2f}")
    if r_mt.ses[i] > 0.6:
        flags.append(f"SE_MT={r_mt.ses[i]:.2f}")
    # Divergência forte CN vs MT (>250 pts) — perfil muito assimétrico
    diff = abs(r_cn.notas[i] - r_mt.notas[i])
    if diff > 250:
        flags.append(f"DIFF_CN_MT={diff:.0f}")
    return "; ".join(flags)


def gerar_alunos_df(r_cn: AreaResult, r_mt: AreaResult) -> pd.DataFrame:
    rows = []
    for i, aluno_id in enumerate(r_cn.ids):
        nota_cn = float(r_cn.notas[i])
        nota_mt = float(r_mt.notas[i])
        media = (nota_cn + nota_mt) / 2
        rows.append({
            "ID": aluno_id,
            "ACERTOS CN": int(r_cn.acertos[i]),
            "θ CN": round(float(r_cn.thetas[i]), 3),
            "SE θ CN": round(float(r_cn.ses[i]), 3),
            "NOTA CN": round(nota_cn, 1),
            "NÍVEL CN": nivel_pedagogico(nota_cn),
            "ACERTOS MT": int(r_mt.acertos[i]),
            "θ MT": round(float(r_mt.thetas[i]), 3),
            "SE θ MT": round(float(r_mt.ses[i]), 3),
            "NOTA MT": round(nota_mt, 1),
            "NÍVEL MT": nivel_pedagogico(nota_mt),
            "MÉDIA CN+MT": round(media, 1),
            "NÍVEL MÉDIO": nivel_pedagogico(media),
            "ATÍPICO": flag_atipico(r_cn, r_mt, i),
        })
    df = pd.DataFrame(rows)
    df["RANK MÉDIA"] = df["MÉDIA CN+MT"].rank(ascending=False, method="min").astype(int)
    df = df.sort_values("RANK MÉDIA").reset_index(drop=True)
    return df


def gerar_por_nivel(df: pd.DataFrame) -> pd.DataFrame:
    """Estatísticas agregadas por nível usando NÍVEL MÉDIO."""
    ordem = ["AVANÇADO", "PROFICIENTE", "BÁSICO", "ELEMENTAR", "INICIANTE"]
    faixas = {"AVANÇADO": "≥650", "PROFICIENTE": "≥580", "BÁSICO": "≥480",
              "ELEMENTAR": "≥380", "INICIANTE": "<380"}
    rows = []
    total = len(df)
    for nivel in ordem:
        sub = df[df["NÍVEL MÉDIO"] == nivel]
        n = len(sub)
        if n == 0:
            rows.append({"NÍVEL": nivel, "FAIXA": faixas[nivel], "N": 0, "%": 0.0,
                         "NOTA CN MED": None, "NOTA MT MED": None, "MÉDIA MED": None,
                         "ACERTOS CN MED": None, "ACERTOS MT MED": None})
            continue
        rows.append({
            "NÍVEL": nivel,
            "FAIXA": faixas[nivel],
            "N": n,
            "%": round(100 * n / total, 1),
            "NOTA CN MED": round(sub["NOTA CN"].median(), 1),
            "NOTA MT MED": round(sub["NOTA MT"].median(), 1),
            "MÉDIA MED": round(sub["MÉDIA CN+MT"].median(), 1),
            "ACERTOS CN MED": int(sub["ACERTOS CN"].median()),
            "ACERTOS MT MED": int(sub["ACERTOS MT"].median()),
        })
    return pd.DataFrame(rows)


def gerar_grafico(df: pd.DataFrame, png_path: Path):
    fig, axes = plt.subplots(1, 3, figsize=(20, 6), dpi=150)
    bins = np.linspace(0, 1000, 51)

    # CN histogram
    axes[0].hist(df["NOTA CN"], bins=bins, color="#2E86AB", edgecolor="white", alpha=0.85)
    axes[0].set_title(f"CN — μ={df['NOTA CN'].mean():.1f}  σ={df['NOTA CN'].std():.1f}",
                       fontsize=13, fontweight="bold")
    axes[0].set_xlabel("Nota CN (escala ENEM 2024)")
    axes[0].set_ylabel("Nº alunos")
    for lim, lab in [(650, "AV"), (580, "PR"), (480, "BA"), (380, "EL")]:
        axes[0].axvline(lim, color="gray", linestyle="--", alpha=0.5)
        axes[0].text(lim, axes[0].get_ylim()[1] * 0.95, lab, fontsize=9,
                     ha="center", color="gray")

    # MT histogram
    axes[1].hist(df["NOTA MT"], bins=bins, color="#A23B72", edgecolor="white", alpha=0.85)
    axes[1].set_title(f"MT — μ={df['NOTA MT'].mean():.1f}  σ={df['NOTA MT'].std():.1f}",
                       fontsize=13, fontweight="bold")
    axes[1].set_xlabel("Nota MT (escala ENEM 2024)")
    axes[1].set_ylabel("Nº alunos")
    for lim, lab in [(650, "AV"), (580, "PR"), (480, "BA"), (380, "EL")]:
        axes[1].axvline(lim, color="gray", linestyle="--", alpha=0.5)
        axes[1].text(lim, axes[1].get_ylim()[1] * 0.95, lab, fontsize=9,
                     ha="center", color="gray")

    # Scatter CN vs MT
    axes[2].scatter(df["NOTA CN"], df["NOTA MT"], s=18, alpha=0.55, color="#264653",
                    edgecolor="white", linewidth=0.4)
    axes[2].plot([0, 1000], [0, 1000], "--", color="gray", linewidth=1, alpha=0.7)
    axes[2].set_xlim(200, 900)
    axes[2].set_ylim(200, 900)
    axes[2].set_xlabel("Nota CN")
    axes[2].set_ylabel("Nota MT")
    corr = df[["NOTA CN", "NOTA MT"]].corr().iloc[0, 1]
    axes[2].set_title(f"CN × MT — ρ={corr:.3f}", fontsize=13, fontweight="bold")
    axes[2].grid(True, alpha=0.3)

    plt.suptitle("FREDÃO — Simulado 2 — V3 + escala σ·θ+μ INEP 2024",
                 fontsize=15, fontweight="bold", y=1.02)
    plt.tight_layout()
    plt.savefig(png_path, dpi=200, bbox_inches="tight")
    plt.close()
    print(f"\n  PNG: {png_path}")


def gerar_resumo(r_cn: AreaResult, r_mt: AreaResult, df: pd.DataFrame, t0: float) -> pd.DataFrame:
    elapsed = time.time() - t0
    corr = df[["NOTA CN", "NOTA MT"]].corr().iloc[0, 1]
    return pd.DataFrame([
        ("Simulado", "Fredão — Simulado 2"),
        ("Data processamento", time.strftime("%Y-%m-%d %H:%M")),
        ("Tempo total", f"{elapsed:.1f}s"),
        ("Alunos processados", r_cn.n_alunos),
        ("Itens CN", r_cn.n_itens_validos),
        ("Itens MT (Q148 anulada excluída)", r_mt.n_itens_validos),
        ("Motor", "V3 cientifico — MML/EM 2PL, c=0.20 fixo"),
        ("Transformação", "σ·θ+μ INEP 2024 (oficial)"),
        ("CN σ", SIGMA_CN),
        ("CN μ", MU_CN),
        ("MT σ", SIGMA_MT),
        ("MT μ", MU_MT),
        ("CN nota μ", f"{df['NOTA CN'].mean():.1f}"),
        ("CN nota σ", f"{df['NOTA CN'].std():.1f}"),
        ("CN nota min/max", f"{df['NOTA CN'].min():.0f} / {df['NOTA CN'].max():.0f}"),
        ("MT nota μ", f"{df['NOTA MT'].mean():.1f}"),
        ("MT nota σ", f"{df['NOTA MT'].std():.1f}"),
        ("MT nota min/max", f"{df['NOTA MT'].min():.0f} / {df['NOTA MT'].max():.0f}"),
        ("Correlação CN × MT", f"{corr:.3f}"),
        ("Atípicos sinalizados", int((df['ATÍPICO'] != '').sum())),
    ], columns=["Métrica", "Valor"])


def gerar_atipicos(df: pd.DataFrame) -> pd.DataFrame:
    sub = df[df["ATÍPICO"] != ""].copy()
    sub = sub.sort_values("MÉDIA CN+MT", ascending=False).reset_index(drop=True)
    return sub


def main():
    t0 = time.time()
    print("FREDÃO — Simulado 2 — pipeline V3 + σ·θ+μ INEP 2024")
    print("="*70)

    r_cn = processar_area("CN", CN_PATH, SIGMA_CN, MU_CN, anulada_idx=-1)
    r_mt = processar_area("MT", MT_PATH, SIGMA_MT, MU_MT, anulada_idx=MT_ANULADA_IDX)

    print("\n=== Consolidando ===")
    df = gerar_alunos_df(r_cn, r_mt)
    df_nivel = gerar_por_nivel(df)
    df_atip = gerar_atipicos(df)
    df_resumo = gerar_resumo(r_cn, r_mt, df, t0)

    ts = time.strftime("%Y%m%d_%H%M%S")
    xlsx_path = OUT_DIR / f"xtri_FREDAO_sim2_{ts}.xlsx"
    png_path = OUT_DIR / f"xtri_FREDAO_sim2_{ts}.png"

    gerar_grafico(df, png_path)

    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        df_resumo.to_excel(writer, sheet_name="Resumo Executivo", index=False)
        df.to_excel(writer, sheet_name="Alunos", index=False)
        df_nivel.to_excel(writer, sheet_name="Por Nivel", index=False)
        df_atip.to_excel(writer, sheet_name="Atipicos - Revisar", index=False)
        # Aba de gráfico
        wb = writer.book
        ws = wb.create_sheet("Grafico")
        from openpyxl.drawing.image import Image as XLImage
        ws.add_image(XLImage(str(png_path)), "B2")

    # Highlight atipicos rows em rosa claro
    wb = openpyxl.load_workbook(xlsx_path)
    ws_al = wb["Alunos"]
    from openpyxl.styles import PatternFill, Font, Alignment
    rosa = PatternFill(start_color="FFE6F0", end_color="FFE6F0", fill_type="solid")
    bold = Font(bold=True)
    centro = Alignment(horizontal="center")
    headers = [c.value for c in ws_al[1]]
    atip_col = headers.index("ATÍPICO") + 1
    for row in range(2, ws_al.max_row + 1):
        if ws_al.cell(row, atip_col).value:
            for col in range(1, ws_al.max_column + 1):
                ws_al.cell(row, col).fill = rosa
    for cell in ws_al[1]:
        cell.font = bold
        cell.alignment = centro
    # ajustar largura
    for col_idx, col in enumerate(ws_al.columns, 1):
        max_len = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws_al.column_dimensions[col[0].column_letter].width = min(max_len + 2, 28)

    # mesma formatação para resumo + nível + atípicos
    for sheet in ["Resumo Executivo", "Por Nivel", "Atipicos - Revisar"]:
        ws = wb[sheet]
        for cell in ws[1]:
            cell.font = bold
            cell.alignment = centro
        for col in ws.columns:
            max_len = max((len(str(c.value)) if c.value is not None else 0) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 38)

    wb.save(xlsx_path)
    print(f"\n  XLSX: {xlsx_path}")
    print(f"\nTOTAL: {time.time()-t0:.1f}s")

    # Quick summary on stdout
    print("\nDistribuição (NÍVEL MÉDIO):")
    print(df_nivel.to_string(index=False))
    print(f"\nAtípicos: {len(df_atip)} alunos")

    return xlsx_path, png_path


if __name__ == "__main__":
    main()
