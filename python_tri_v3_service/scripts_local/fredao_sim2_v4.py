"""Fredão Sim2 — Análise V4 INEP OFICIAL (params autoritários do INEP).

Diferente das corridas anteriores (V3 local), esta:
- Usa parâmetros a, b, c OFICIAIS INEP para cada item por (ano, CO_PROVA_AZUL, posição).
- 2024: V4 PURO (a,b,c + σ·θ+μ oficiais).
- 2017: V4 híbrido (a,b,c oficiais INEP + ancoragem por tabela pedagógica 2017,
  porque σ/μ INEP 2017 não foi pré-calculado).

CO_PROVAs AZUL regular:
- 2017 CN=391, MT=403  (gabarito oficial INEP 2017)
- 2024 CN=1333, MT=1323 (gabarito oficial INEP 2024 1ª aplicação)

Q148 do Fredão (anulada pelo cursinho) — excluída do EAP em todas as rodadas.
"""

from __future__ import annotations

import json
import sys
import time
from pathlib import Path
from typing import Dict, List, Tuple

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from app.irt.scoring import eap_theta  # noqa: E402
from app.irt.transform import ancorar_notas  # noqa: E402
from app.tabela_referencia import TabelaReferenciaTRI  # noqa: E402

DROPBOX = Path("/Users/home/Library/CloudStorage/Dropbox/Fredao Simulados")
CN_PATH = DROPBOX / "Natureza - Simulado 2.xlsx"
MT_PATH = DROPBOX / "Matemática - Simulado 2.xlsx"

PARAMS_JSON = Path(__file__).parent / "data" / "params_oficiais.json"

MT_ANULADA_IDX = 12  # Q148 do Fredão (posição 13 do bloco MT, índice 12)

# σ, μ oficiais INEP 2024 (de enem_transformacoes)
SIGMA_CN_2024, MU_CN_2024 = 113.1042, 501.1432
SIGMA_MT_2024, MU_MT_2024 = 129.6484, 500.0186

OUT_DIR = Path("/Users/home/Downloads/XTRI Engine")
OUT_DIR.mkdir(parents=True, exist_ok=True)


def carregar(path: Path) -> Tuple[List[str], np.ndarray]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    ids = []
    rows = []
    for r in range(2, ws.max_row + 1):
        aluno = ws.cell(r, 1).value
        if not aluno:
            continue
        ids.append(str(aluno))
        rows.append([int(ws.cell(r, c).value or 0) for c in range(2, 47)])
    return ids, np.array(rows, dtype=int)


def carregar_params() -> Dict:
    with open(PARAMS_JSON) as f:
        return json.load(f)


def params_para_dicts(params_list: List[List]) -> List[Dict]:
    """Converte [[posicao, a, b, c, gab], ...] → [{a, b, c, ...}, ...]"""
    return [{"a": float(p[1]), "b": float(p[2]), "c": float(p[3]),
             "posicao": int(p[0]), "gabarito": p[4]} for p in params_list]


def nivel_pedagogico(nota: float) -> str:
    if nota >= 650:
        return "AVANÇADO"
    if nota >= 580:
        return "PROFICIENTE"
    if nota >= 480:
        return "BÁSICO"
    if nota >= 380:
        return "ELEMENTAR"
    return "INICIANTE"


def estimar_thetas(resp_matrix: np.ndarray, params: List[Dict]) -> Tuple[np.ndarray, np.ndarray]:
    n = len(resp_matrix)
    thetas = np.zeros(n)
    ses = np.zeros(n)
    for i, row in enumerate(resp_matrix):
        t, se = eap_theta(row, params)
        thetas[i] = t
        ses[i] = se
    return thetas, ses


def ancorar_tabela(area: str, thetas: np.ndarray, acertos: np.ndarray,
                   ses: np.ndarray, tab: TabelaReferenciaTRI) -> np.ndarray:
    thetas_list = []
    for i in range(len(thetas)):
        thetas_list.append({
            "aluno_id": f"row_{i}", "nome": f"row_{i}",
            "acertos": int(acertos[i]),
            "theta": float(thetas[i]), "se": float(ses[i]),
            "acertos_faceis": 0, "acertos_medios": 0,
            "acertos_dificeis": 0, "acert_dif": 0.0,
        })
    res = ancorar_notas(thetas_list, tab, area)
    return np.array([r["nota_ancorada"] for r in res])


def aplicar_sigma_mu(thetas: np.ndarray, sigma: float, mu: float) -> np.ndarray:
    return np.clip(sigma * thetas + mu, 0.0, 1000.0)


def processar_area_ano(area: str, ano: int, resp_matrix: np.ndarray,
                        params_oficiais: Dict,
                        tabelas: Dict[int, TabelaReferenciaTRI],
                        excluir_idx: List[int] = None) -> Dict:
    """Processa uma combinação área × ano com V4."""
    print(f"\n  {area} {ano}:")
    params_list = params_oficiais[str(ano)][area]
    params = params_para_dicts(params_list)

    # Excluir Q148 (idx 12 do MT) se aplicável
    if excluir_idx:
        # Excluir tanto dos params quanto das respostas para alinhamento
        idx_keep = [j for j in range(len(params)) if j not in excluir_idx]
        params = [params[j] for j in idx_keep]
        resp = resp_matrix[:, idx_keep]
    else:
        resp = resp_matrix

    a_vals = [p["a"] for p in params]
    b_vals = [p["b"] for p in params]
    c_vals = [p["c"] for p in params]
    print(f"    Params INEP: {len(params)} itens, "
          f"a μ={np.mean(a_vals):.2f}, b μ={np.mean(b_vals):.2f}, c μ={np.mean(c_vals):.2f}")

    thetas, ses = estimar_thetas(resp, params)
    acertos = resp.sum(axis=1)
    print(f"    θ μ={thetas.mean():+.3f} σ={thetas.std():.3f}")

    # Conversão θ → nota
    if ano == 2024:
        # V4 puro: σ·θ+μ oficial
        sigma = SIGMA_CN_2024 if area == "CN" else SIGMA_MT_2024
        mu = MU_CN_2024 if area == "CN" else MU_MT_2024
        notas = aplicar_sigma_mu(thetas, sigma, mu)
        metodo = f"σ·θ+μ INEP 2024 ({sigma:.4f}·θ + {mu:.4f})"
    else:
        # Híbrido: tabela 2017
        notas = ancorar_tabela(area, thetas, acertos, ses, tabelas[ano])
        metodo = "tabela_pedagogica INEP 2017 (tri_min/tri_max)"
    print(f"    {metodo}")
    print(f"    Notas μ={notas.mean():.1f} σ={notas.std():.1f} "
          f"min={notas.min():.0f} max={notas.max():.0f}")

    return {"thetas": thetas, "ses": ses, "acertos": acertos,
            "notas": notas, "metodo": metodo, "n_itens": len(params)}


def gerar_grafico_consolidado(df: pd.DataFrame, ano: int, png_path: Path):
    fig, axes = plt.subplots(2, 2, figsize=(16, 12), dpi=150)
    bins = np.linspace(0, 1000, 51)

    cn_col = f"NOTA CN {ano} V4"
    mt_col = f"NOTA MT {ano} V4"
    med_col = f"MÉDIA {ano} V4"

    ax = axes[0, 0]
    ax.hist(df[cn_col], bins=bins, color="#2E86AB", edgecolor="white", alpha=0.85)
    ax.set_title(f"CN V4 INEP {ano} — μ={df[cn_col].mean():.1f}  σ={df[cn_col].std():.1f}",
                 fontsize=13, fontweight="bold")
    ax.set_xlabel(f"Nota CN")
    ax.set_ylabel("Nº alunos")
    for lim, lab in [(650, "AV"), (580, "PR"), (480, "BA"), (380, "EL")]:
        ax.axvline(lim, color="gray", linestyle="--", alpha=0.5)

    ax = axes[0, 1]
    ax.hist(df[mt_col], bins=bins, color="#A23B72", edgecolor="white", alpha=0.85)
    ax.set_title(f"MT V4 INEP {ano} — μ={df[mt_col].mean():.1f}  σ={df[mt_col].std():.1f}",
                 fontsize=13, fontweight="bold")
    ax.set_xlabel(f"Nota MT")
    ax.set_ylabel("Nº alunos")
    for lim, lab in [(650, "AV"), (580, "PR"), (480, "BA"), (380, "EL")]:
        ax.axvline(lim, color="gray", linestyle="--", alpha=0.5)

    ax = axes[1, 0]
    ax.hist(df[med_col], bins=bins, color="#F18F01", edgecolor="white", alpha=0.85)
    ax.set_title(f"MÉDIA CN+MT V4 INEP {ano} — μ={df[med_col].mean():.1f}  σ={df[med_col].std():.1f}",
                 fontsize=13, fontweight="bold")
    ax.set_xlabel("Nota média")
    ax.set_ylabel("Nº alunos")
    for lim, lab in [(650, "AV"), (580, "PR"), (480, "BA"), (380, "EL")]:
        ax.axvline(lim, color="gray", linestyle="--", alpha=0.5)

    ax = axes[1, 1]
    ax.scatter(df[cn_col], df[mt_col], s=18, alpha=0.55,
               color="#264653", edgecolor="white", linewidth=0.4)
    ax.plot([0, 1000], [0, 1000], "--", color="gray", linewidth=1, alpha=0.7)
    lim_min = min(df[cn_col].min(), df[mt_col].min()) - 20
    lim_max = max(df[cn_col].max(), df[mt_col].max()) + 20
    ax.set_xlim(lim_min, lim_max)
    ax.set_ylim(lim_min, lim_max)
    ax.set_xlabel("Nota CN V4")
    ax.set_ylabel("Nota MT V4")
    corr = df[[cn_col, mt_col]].corr().iloc[0, 1]
    ax.set_title(f"CN × MT V4 — ρ={corr:.3f}", fontsize=13, fontweight="bold")
    ax.grid(True, alpha=0.3)

    plt.suptitle(f"FREDÃO Simulado 2 — V4 INEP OFICIAL — escala {ano}",
                 fontsize=16, fontweight="bold", y=1.00)
    plt.tight_layout()
    plt.savefig(png_path, dpi=200, bbox_inches="tight")
    plt.close()


def main():
    t0 = time.time()
    print("FREDÃO Sim2 — V4 INEP OFICIAL")
    print("=" * 70)

    ids_cn, mat_cn = carregar(CN_PATH)
    ids_mt, mat_mt = carregar(MT_PATH)
    if ids_cn != ids_mt:
        raise RuntimeError("IDs CN ≠ MT")
    ids = ids_cn
    print(f"  Alunos: {len(ids)}")
    print(f"  CN: {mat_cn.shape[1]} itens | MT: {mat_mt.shape[1]} itens (Q148 anulada)")

    params_oficiais = carregar_params()
    print(f"  Params oficiais carregados: {PARAMS_JSON.name}")
    print(f"  Cobertura: 2017 CN=391, MT=403 | 2024 CN=1333, MT=1323 (AZUL regular)")

    tabelas = TabelaReferenciaTRI.carregar_multiano()
    print(f"  Tabela multianos: {sorted(tabelas.keys())}")

    # Processar cada combinação
    resultados = {}
    for ano in [2017, 2024]:
        # CN: 45 itens, sem exclusões
        resultados[("CN", ano)] = processar_area_ano(
            "CN", ano, mat_cn, params_oficiais, tabelas, excluir_idx=None,
        )
        # MT: 45 itens, excluir Q148
        resultados[("MT", ano)] = processar_area_ano(
            "MT", ano, mat_mt, params_oficiais, tabelas, excluir_idx=[MT_ANULADA_IDX],
        )

    # === Montar DataFrame ===
    rows = []
    for i, aluno_id in enumerate(ids):
        row = {"ID": aluno_id}
        for area in ["CN", "MT"]:
            row[f"ACERTOS {area}"] = int(resultados[(area, 2024)]["acertos"][i])
        for ano in [2017, 2024]:
            for area in ["CN", "MT"]:
                r = resultados[(area, ano)]
                row[f"θ {area} {ano} V4"] = round(float(r["thetas"][i]), 3)
                row[f"SE θ {area} {ano} V4"] = round(float(r["ses"][i]), 3)
                row[f"NOTA {area} {ano} V4"] = round(float(r["notas"][i]), 1)
                row[f"NÍVEL {area} {ano} V4"] = nivel_pedagogico(float(r["notas"][i]))
            ncn = float(resultados[("CN", ano)]["notas"][i])
            nmt = float(resultados[("MT", ano)]["notas"][i])
            med = (ncn + nmt) / 2
            row[f"MÉDIA {ano} V4"] = round(med, 1)
            row[f"NÍVEL MÉDIO {ano} V4"] = nivel_pedagogico(med)
        row["DIFF MÉDIA 2024-2017 V4"] = round(row["MÉDIA 2024 V4"] - row["MÉDIA 2017 V4"], 1)
        # Atípicos
        flags = []
        for ano in [2017, 2024]:
            for area in ["CN", "MT"]:
                t = resultados[(area, ano)]["thetas"][i]
                se = resultados[(area, ano)]["ses"][i]
                if abs(t) > 2.5:
                    flags.append(f"θ_{area}_{ano}={t:+.2f}")
                if se > 0.6:
                    flags.append(f"SE_{area}_{ano}={se:.2f}")
        row["ATÍPICO"] = "; ".join(flags)
        rows.append(row)

    df = pd.DataFrame(rows)
    df["RANK MÉDIA 2024 V4"] = df["MÉDIA 2024 V4"].rank(ascending=False, method="min").astype(int)
    df = df.sort_values("RANK MÉDIA 2024 V4").reset_index(drop=True)

    # === Por Nível por ano ===
    def por_nivel(df: pd.DataFrame, ano: int) -> pd.DataFrame:
        col_n = f"NÍVEL MÉDIO {ano} V4"
        ordem = ["AVANÇADO", "PROFICIENTE", "BÁSICO", "ELEMENTAR", "INICIANTE"]
        faixas = {"AVANÇADO": "≥650", "PROFICIENTE": "≥580", "BÁSICO": "≥480",
                  "ELEMENTAR": "≥380", "INICIANTE": "<380"}
        rows_n = []
        total = len(df)
        for nv in ordem:
            sub = df[df[col_n] == nv]
            n = len(sub)
            if n == 0:
                rows_n.append({"NÍVEL": nv, "FAIXA": faixas[nv], "N": 0, "%": 0.0,
                               "CN MED": None, "MT MED": None, "MÉDIA MED": None})
                continue
            rows_n.append({
                "NÍVEL": nv,
                "FAIXA": faixas[nv],
                "N": n,
                "%": round(100 * n / total, 1),
                "CN MED": round(sub[f"NOTA CN {ano} V4"].median(), 1),
                "MT MED": round(sub[f"NOTA MT {ano} V4"].median(), 1),
                "MÉDIA MED": round(sub[f"MÉDIA {ano} V4"].median(), 1),
                "ACERTOS CN MED": int(sub["ACERTOS CN"].median()),
                "ACERTOS MT MED": int(sub["ACERTOS MT"].median()),
            })
        return pd.DataFrame(rows_n)

    df_niv_2017 = por_nivel(df, 2017)
    df_niv_2024 = por_nivel(df, 2024)

    ord_full = ["AVANÇADO", "PROFICIENTE", "BÁSICO", "ELEMENTAR", "INICIANTE"]
    comparacao = pd.crosstab(
        df["NÍVEL MÉDIO 2017 V4"], df["NÍVEL MÉDIO 2024 V4"],
        rownames=["Nível 2017 ↓"], colnames=["Nível 2024 →"], dropna=False,
    ).reindex(index=ord_full, columns=ord_full, fill_value=0)

    elapsed_total = time.time() - t0
    n_atip = int((df["ATÍPICO"] != "").sum())
    resumo_rows = [
        ("Simulado", "Fredão — Simulado 2"),
        ("Análise", "V4 INEP OFICIAL (params autoritativos)"),
        ("Alunos", len(ids)),
        ("Itens CN (oficiais)", "45"),
        ("Itens MT (oficiais)", "44 (Q148 Fredão excluída)"),
        ("CO_PROVA AZUL 2017 CN", 391),
        ("CO_PROVA AZUL 2017 MT", 403),
        ("CO_PROVA AZUL 2024 CN", 1333),
        ("CO_PROVA AZUL 2024 MT", 1323),
        ("Fonte params", "Supabase enem_itens + enem_itens_prova"),
        ("Tempo total", f"{elapsed_total:.1f}s"),
        ("", ""),
        ("── Escala 2017 (V4 híbrido) ──", ""),
        ("Método θ", "EAP com a,b,c oficiais INEP 2017"),
        ("Método nota", "Ancoragem por tabela pedagógica 2017"),
        ("CN nota μ ± σ",
            f"{df['NOTA CN 2017 V4'].mean():.1f} ± {df['NOTA CN 2017 V4'].std():.1f}"),
        ("MT nota μ ± σ",
            f"{df['NOTA MT 2017 V4'].mean():.1f} ± {df['NOTA MT 2017 V4'].std():.1f}"),
        ("MÉDIA μ ± σ",
            f"{df['MÉDIA 2017 V4'].mean():.1f} ± {df['MÉDIA 2017 V4'].std():.1f}"),
        ("Correlação CN × MT 2017",
            f"{df[['NOTA CN 2017 V4', 'NOTA MT 2017 V4']].corr().iloc[0,1]:.3f}"),
        ("", ""),
        ("── Escala 2024 (V4 puro) ──", ""),
        ("Método θ", "EAP com a,b,c oficiais INEP 2024"),
        ("Método nota", f"σ·θ+μ oficial (CN: {SIGMA_CN_2024}·θ+{MU_CN_2024}, "
                          f"MT: {SIGMA_MT_2024}·θ+{MU_MT_2024})"),
        ("CN nota μ ± σ",
            f"{df['NOTA CN 2024 V4'].mean():.1f} ± {df['NOTA CN 2024 V4'].std():.1f}"),
        ("MT nota μ ± σ",
            f"{df['NOTA MT 2024 V4'].mean():.1f} ± {df['NOTA MT 2024 V4'].std():.1f}"),
        ("MÉDIA μ ± σ",
            f"{df['MÉDIA 2024 V4'].mean():.1f} ± {df['MÉDIA 2024 V4'].std():.1f}"),
        ("Correlação CN × MT 2024",
            f"{df[['NOTA CN 2024 V4', 'NOTA MT 2024 V4']].corr().iloc[0,1]:.3f}"),
        ("", ""),
        ("── Comparação 2017 vs 2024 ──", ""),
        ("DIFF MÉDIA 2024-2017 μ", f"{df['DIFF MÉDIA 2024-2017 V4'].mean():+.1f}"),
        ("DIFF MÉDIA 2024-2017 σ", f"{df['DIFF MÉDIA 2024-2017 V4'].std():.1f}"),
        ("DIFF min/max", f"{df['DIFF MÉDIA 2024-2017 V4'].min():+.1f} / "
                          f"{df['DIFF MÉDIA 2024-2017 V4'].max():+.1f}"),
        ("Atípicos sinalizados", n_atip),
    ]
    resumo = pd.DataFrame(resumo_rows, columns=["Métrica", "Valor"])

    df_atip = df[df["ATÍPICO"] != ""].copy().reset_index(drop=True)

    # === Gráficos ===
    ts = time.strftime("%Y%m%d_%H%M%S")
    png_2017 = OUT_DIR / f"xtri_FREDAO_sim2_V4_escala2017_{ts}.png"
    png_2024 = OUT_DIR / f"xtri_FREDAO_sim2_V4_escala2024_{ts}.png"
    gerar_grafico_consolidado(df, 2017, png_2017)
    gerar_grafico_consolidado(df, 2024, png_2024)
    print(f"\n  PNG 2017: {png_2017.name}")
    print(f"  PNG 2024: {png_2024.name}")

    # === Excel ===
    xlsx_path = OUT_DIR / f"xtri_FREDAO_sim2_V4_{ts}.xlsx"
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        resumo.to_excel(writer, sheet_name="Resumo Executivo", index=False)
        df.to_excel(writer, sheet_name="Alunos", index=False)
        df_niv_2017.to_excel(writer, sheet_name="Por Nivel 2017", index=False)
        df_niv_2024.to_excel(writer, sheet_name="Por Nivel 2024", index=False)
        comparacao.to_excel(writer, sheet_name="Migracao 2017-2024")
        df_atip.to_excel(writer, sheet_name="Atipicos - Revisar", index=False)
        wb = writer.book
        from openpyxl.drawing.image import Image as XLImage
        ws_g17 = wb.create_sheet("Grafico 2017")
        ws_g17.add_image(XLImage(str(png_2017)), "B2")
        ws_g24 = wb.create_sheet("Grafico 2024")
        ws_g24.add_image(XLImage(str(png_2024)), "B2")

    # Formatação
    wb = openpyxl.load_workbook(xlsx_path)
    from openpyxl.styles import PatternFill, Font, Alignment
    rosa = PatternFill(start_color="FFE6F0", end_color="FFE6F0", fill_type="solid")
    bold = Font(bold=True)
    centro = Alignment(horizontal="center")

    ws_al = wb["Alunos"]
    headers = [c.value for c in ws_al[1]]
    atip_col = headers.index("ATÍPICO") + 1
    for r in range(2, ws_al.max_row + 1):
        if ws_al.cell(r, atip_col).value:
            for c in range(1, ws_al.max_column + 1):
                ws_al.cell(r, c).fill = rosa

    for sheet in ["Resumo Executivo", "Alunos", "Por Nivel 2017",
                  "Por Nivel 2024", "Migracao 2017-2024", "Atipicos - Revisar"]:
        ws = wb[sheet]
        for cell in ws[1]:
            cell.font = bold
            cell.alignment = centro
        for col in ws.columns:
            max_len = max((len(str(c.value)) if c.value is not None else 0) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 32)

    wb.save(xlsx_path)
    print(f"  XLSX: {xlsx_path.name}")
    print(f"\nTOTAL: {elapsed_total:.1f}s")

    print("\n── Por Nível MÉDIO escala 2017 (V4 híbrido) ──")
    print(df_niv_2017.to_string(index=False))
    print("\n── Por Nível MÉDIO escala 2024 (V4 puro) ──")
    print(df_niv_2024.to_string(index=False))
    print("\n── Migração ──")
    print(comparacao.to_string())
    print(f"\nAtípicos: {n_atip}")
    print(f"DIFF MÉDIA 2024-2017: μ={df['DIFF MÉDIA 2024-2017 V4'].mean():+.1f}, "
          f"min={df['DIFF MÉDIA 2024-2017 V4'].min():+.1f}, "
          f"max={df['DIFF MÉDIA 2024-2017 V4'].max():+.1f}")


if __name__ == "__main__":
    main()
