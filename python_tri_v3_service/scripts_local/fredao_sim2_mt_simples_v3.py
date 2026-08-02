"""Fredão Sim2 — MT relatório 4 colunas — V3 mapeamento TCC↔tabela.

Método robusto sem saturação:
  E[acertos | θ] = TCC(θ) com params oficiais INEP do ano
  nota = interpolação contínua entre tri_med[floor(E)] e tri_med[ceil(E)]

Cada θ → único E[acertos] → única nota. Não satura por banda de acertos,
e usa a tabela INEP autoritativamente (não regressão linear).

Para validação, também roda 2024 via σ·θ+μ oficial e compara.
"""

from __future__ import annotations

import json
import sys
import time
from pathlib import Path
from typing import Dict, List, Tuple

import numpy as np
import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from app.irt.scoring import eap_theta  # noqa: E402
from app.tabela_referencia import (  # noqa: E402
    TRI_MAXIMA_OFICIAL, TRI_MINIMA_OFICIAL, TabelaReferenciaTRI,
)

MT_PATH = Path("/Users/home/Library/CloudStorage/Dropbox/Fredao Simulados/Matemática - Simulado 2.xlsx")
PARAMS_JSON = Path(__file__).parent / "data" / "params_oficiais.json"
MT_ANULADA_IDX = 12

SIGMA_MT_2024_OFICIAL, MU_MT_2024_OFICIAL = 129.6484, 500.0186

OUT_DIR = Path("/Users/home/Downloads/XTRI Engine")


def prob_acerto_3pl(theta: float, a: float, b: float, c: float) -> float:
    return c + (1.0 - c) / (1.0 + np.exp(-a * (theta - b)))


def tcc_acertos_esperados(theta: float, params: List[Dict]) -> float:
    return float(sum(prob_acerto_3pl(theta, p["a"], p["b"], p["c"]) for p in params))


def nota_via_tcc_tabela(theta: float, params: List[Dict],
                         tabela: TabelaReferenciaTRI, area: str, n_itens: int) -> float:
    """θ → E[acertos|θ] → interpolação entre tri_med adjacentes."""
    e_ac = tcc_acertos_esperados(theta, params)
    e_ac = max(0.0, min(float(n_itens), e_ac))
    k_lo = int(np.floor(e_ac))
    k_hi = min(n_itens, k_lo + 1)
    frac = (e_ac - k_lo) if k_hi > k_lo else 0.0
    nota_lo = tabela.obter(area, k_lo)["tri_med"]
    nota_hi = tabela.obter(area, k_hi)["tri_med"]
    return nota_lo + frac * (nota_hi - nota_lo)


def carregar_mt() -> Tuple[List[str], np.ndarray]:
    wb = openpyxl.load_workbook(MT_PATH, data_only=True)
    ws = wb.worksheets[0]
    ids, rows = [], []
    for r in range(2, ws.max_row + 1):
        aluno = ws.cell(r, 1).value
        if not aluno:
            continue
        ids.append(str(aluno))
        rows.append([int(ws.cell(r, c).value or 0) for c in range(2, 47)])
    return ids, np.array(rows, dtype=int)


def params_para_dicts(plist: List[List]) -> List[Dict]:
    return [{"a": float(p[1]), "b": float(p[2]), "c": float(p[3])} for p in plist]


def estimar_thetas(resp: np.ndarray, params: List[Dict]) -> np.ndarray:
    n = len(resp)
    thetas = np.zeros(n)
    for i, row in enumerate(resp):
        t, _ = eap_theta(row, params)
        thetas[i] = t
    return thetas


def main():
    t0 = time.time()
    print("FREDÃO Sim2 — MT relatório SIMPLES V3 (TCC↔tabela)")
    print("=" * 70)

    ids, mat_full = carregar_mt()
    keep = [j for j in range(mat_full.shape[1]) if j != MT_ANULADA_IDX]
    mat = mat_full[:, keep]
    acertos = mat.sum(axis=1).astype(int)
    print(f"  Alunos: {len(ids)} | itens válidos: {mat.shape[1]} (Q148 excluída)")

    with open(PARAMS_JSON) as f:
        params_of = json.load(f)

    params_2017_full = params_para_dicts(params_of["2017"]["MT"])
    params_2024_full = params_para_dicts(params_of["2024"]["MT"])
    params_2017 = [params_2017_full[j] for j in keep]
    params_2024 = [params_2024_full[j] for j in keep]

    tabelas = TabelaReferenciaTRI.carregar_multiano()
    n_itens = mat.shape[1]
    tri_min_oficial = TRI_MINIMA_OFICIAL["MT"]
    tri_max_oficial = TRI_MAXIMA_OFICIAL["MT"]

    # Estimar θs
    print("\n  Estimando θ com params INEP 2017...")
    thetas_2017 = estimar_thetas(mat, params_2017)
    print(f"    θ μ={thetas_2017.mean():+.3f} σ={thetas_2017.std():.3f} "
          f"min={thetas_2017.min():+.2f} max={thetas_2017.max():+.2f}")

    print("  Estimando θ com params INEP 2024...")
    thetas_2024 = estimar_thetas(mat, params_2024)
    print(f"    θ μ={thetas_2024.mean():+.3f} σ={thetas_2024.std():.3f} "
          f"min={thetas_2024.min():+.2f} max={thetas_2024.max():+.2f}")

    # ── TRI 2017 via TCC↔tabela ──
    print("\n  Calculando TRI 2017 via TCC↔tabela...")
    notas_2017 = np.array([
        nota_via_tcc_tabela(t, params_2017, tabelas[2017], "MT", n_itens)
        for t in thetas_2017
    ])
    notas_2017 = np.clip(notas_2017, tri_min_oficial, 1000.0)

    # ── TRI 2024 (duas versões pra validação) ──
    print("  Calculando TRI 2024 via σ·θ+μ oficial...")
    notas_2024_oficial = np.clip(
        SIGMA_MT_2024_OFICIAL * thetas_2024 + MU_MT_2024_OFICIAL,
        tri_min_oficial, tri_max_oficial,
    )

    print("  Calculando TRI 2024 via TCC↔tabela (validação)...")
    notas_2024_tcc = np.array([
        nota_via_tcc_tabela(t, params_2024, tabelas[2024], "MT", n_itens)
        for t in thetas_2024
    ])
    notas_2024_tcc = np.clip(notas_2024_tcc, tri_min_oficial, 1000.0)

    # Comparar os dois métodos em 2024
    diff = notas_2024_oficial - notas_2024_tcc
    print(f"\n  2024 σ·θ+μ oficial vs TCC↔tabela:")
    print(f"    μ oficial = {notas_2024_oficial.mean():.1f}, μ TCC = {notas_2024_tcc.mean():.1f}")
    print(f"    σ oficial = {notas_2024_oficial.std():.1f}, σ TCC = {notas_2024_tcc.std():.1f}")
    print(f"    MAE entre métodos: {np.abs(diff).mean():.1f}")
    print(f"    Correlação: {np.corrcoef(notas_2024_oficial, notas_2024_tcc)[0,1]:.4f}")

    # Para o relatório final: usar σ·θ+μ oficial em 2024 (gold standard)
    notas_2024 = notas_2024_oficial

    print(f"\n  TRI 2017 (final): μ={notas_2017.mean():.1f} σ={notas_2017.std():.1f} "
          f"min={notas_2017.min():.0f} max={notas_2017.max():.0f}")
    print(f"  TRI 2024 (final): μ={notas_2024.mean():.1f} σ={notas_2024.std():.1f} "
          f"min={notas_2024.min():.0f} max={notas_2024.max():.0f}")

    # ── Sanity check: notas distintas por mesmo acertos ──
    print("\n  Sanity check — notas distintas para mesmo acertos:")
    df_check = pd.DataFrame({"acertos": acertos, "TRI 2017": notas_2017, "TRI 2024": notas_2024})
    for k in [38, 39, 40, 41, 42, 43, 0, 1, 2]:
        sub = df_check[df_check.acertos == k]
        if len(sub) > 1:
            print(f"    {k:2d} acertos ({len(sub):2d} alunos): "
                  f"2017 distintas={sub['TRI 2017'].nunique():2d}, "
                  f"2024 distintas={sub['TRI 2024'].nunique():2d}, "
                  f"2017 range=[{sub['TRI 2017'].min():.1f}, {sub['TRI 2017'].max():.1f}], "
                  f"2024 range=[{sub['TRI 2024'].min():.1f}, {sub['TRI 2024'].max():.1f}]")

    # ── DataFrame final ──
    df = pd.DataFrame({
        "Nome": ids,
        "Acertos": acertos,
        "TRI 2017": np.round(notas_2017, 1),
        "TRI 2024": np.round(notas_2024, 1),
    })
    df = df.sort_values("TRI 2024", ascending=False).reset_index(drop=True)

    ts = time.strftime("%Y%m%d_%H%M%S")
    xlsx_path = OUT_DIR / f"xtri_FREDAO_sim2_MT_simples_{ts}.xlsx"
    df.to_excel(xlsx_path, sheet_name="MT", index=False)

    wb = openpyxl.load_workbook(xlsx_path)
    from openpyxl.styles import Font, Alignment, PatternFill
    ws = wb["MT"]
    bold = Font(bold=True, color="FFFFFF")
    fill = PatternFill(start_color="2E86AB", end_color="2E86AB", fill_type="solid")
    centro = Alignment(horizontal="center", vertical="center")
    for cell in ws[1]:
        cell.font = bold
        cell.alignment = centro
        cell.fill = fill
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 14
    for r in range(2, ws.max_row + 1):
        for c in range(1, 5):
            ws.cell(r, c).alignment = centro
    ws.freeze_panes = "A2"
    wb.save(xlsx_path)

    print(f"\n  XLSX: {xlsx_path}")
    print(f"\nTOTAL: {time.time()-t0:.1f}s")

    print("\n── Top 10 ──")
    print(df.head(10).to_string(index=False))
    print("\n── Bottom 5 ──")
    print(df.tail(5).to_string(index=False))


if __name__ == "__main__":
    main()
