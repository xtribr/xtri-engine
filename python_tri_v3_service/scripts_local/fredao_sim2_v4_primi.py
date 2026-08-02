"""Fredão Sim2 — V4 INEP OFICIAL com gráficos estilo PRIMI.

Substitui os scatters anteriores por:
- Painel 1 (CN): acertos × nota V4 com BANDA INEP (tri_min/tri_med/tri_max do ano)
  sobreposta como referência pedagógica.
- Painel 2 (MT): idem.
- Painel 3 (CN): acert_dif PRIMI (% itens DIFÍCEIS INEP acertados, b>1.5) × nota V4.
- Painel 4 (MT): idem.

Itens difíceis classificados via b INEP oficial do ano correspondente:
- Fácil: b < -1.5  |  Médio: -1.5 ≤ b ≤ 1.5  |  Difícil: b > 1.5
(Primi & Cicchetto methodology, já documentada em app/motores/cientifico.py)
"""

from __future__ import annotations

import json
import sys
import time
from pathlib import Path
from typing import Dict, List, Tuple

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
import numpy as np
import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from app.irt.scoring import eap_theta  # noqa: E402
from app.irt.transform import ancorar_notas  # noqa: E402
from app.tabela_referencia import TabelaReferenciaTRI  # noqa: E402

DROPBOX = Path("/Users/home/Library/CloudStorage/Dropbox/Fredao Simulados")
CN_PATH = DROPBOX / "Natureza - Simulado 2.xlsx"
MT_PATH = DROPBOX / "Matemática - Simulado 2.xlsx"

PARAMS_JSON = Path(__file__).parent / "data" / "params_oficiais.json"

MT_ANULADA_IDX = 12

SIGMA_CN_2024, MU_CN_2024 = 113.1042, 501.1432
SIGMA_MT_2024, MU_MT_2024 = 129.6484, 500.0186

# Cores dos níveis pedagógicos (em escala de verde→vermelho)
NIVEIS_CORES = {
    "AVANÇADO": "#1B7340",
    "PROFICIENTE": "#56AE57",
    "BÁSICO": "#F5C84C",
    "ELEMENTAR": "#E58A3D",
    "INICIANTE": "#C1392B",
}

OUT_DIR = Path("/Users/home/Downloads/XTRI Engine")
OUT_DIR.mkdir(parents=True, exist_ok=True)


def carregar(path: Path) -> Tuple[List[str], np.ndarray]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    ids, rows = [], []
    for r in range(2, ws.max_row + 1):
        aluno = ws.cell(r, 1).value
        if not aluno:
            continue
        ids.append(str(aluno))
        rows.append([int(ws.cell(r, c).value or 0) for c in range(2, 47)])
    return ids, np.array(rows, dtype=int)


def carregar_params() -> Dict:
    with open(PARAMS_JSON) as f:
        return json.load(f)


def params_para_dicts(params_list: List[List]) -> List[Dict]:
    return [{"a": float(p[1]), "b": float(p[2]), "c": float(p[3]),
             "posicao": int(p[0]), "gabarito": p[4]} for p in params_list]


def nivel_pedagogico(nota: float) -> str:
    if nota >= 650: return "AVANÇADO"
    if nota >= 580: return "PROFICIENTE"
    if nota >= 480: return "BÁSICO"
    if nota >= 380: return "ELEMENTAR"
    return "INICIANTE"


def classificar_itens_primi(params: List[Dict]) -> Tuple[List[int], List[int], List[int]]:
    """Fácil b<-1.5 | Médio -1.5≤b≤1.5 | Difícil b>1.5. Primi & Cicchetto."""
    idx_f = [j for j, p in enumerate(params) if p["b"] < -1.5]
    idx_m = [j for j, p in enumerate(params) if -1.5 <= p["b"] <= 1.5]
    idx_d = [j for j, p in enumerate(params) if p["b"] > 1.5]
    return idx_f, idx_m, idx_d


def computar_primi(resp: np.ndarray, params: List[Dict]) -> Dict:
    idx_f, idx_m, idx_d = classificar_itens_primi(params)
    n = len(resp)
    af = resp[:, idx_f].sum(axis=1) if idx_f else np.zeros(n, dtype=int)
    am = resp[:, idx_m].sum(axis=1) if idx_m else np.zeros(n, dtype=int)
    ad = resp[:, idx_d].sum(axis=1) if idx_d else np.zeros(n, dtype=int)
    nf, nm, nd = len(idx_f), len(idx_m), len(idx_d)
    return {
        "n_faceis": nf, "n_medios": nm, "n_dificeis": nd,
        "acertos_faceis": af, "acertos_medios": am, "acertos_dificeis": ad,
        "acert_dif": (ad / max(1, nd)) if nd > 0 else np.zeros(n),
        "acert_fac": (af / max(1, nf)) if nf > 0 else np.zeros(n),
    }


def estimar_thetas(resp: np.ndarray, params: List[Dict]) -> Tuple[np.ndarray, np.ndarray]:
    n = len(resp)
    thetas = np.zeros(n)
    ses = np.zeros(n)
    for i, row in enumerate(resp):
        t, se = eap_theta(row, params)
        thetas[i] = t
        ses[i] = se
    return thetas, ses


def ancorar_tabela(area: str, thetas: np.ndarray, acertos: np.ndarray,
                   ses: np.ndarray, tab: TabelaReferenciaTRI) -> np.ndarray:
    thetas_list = []
    for i in range(len(thetas)):
        thetas_list.append({
            "aluno_id": f"r{i}", "nome": f"r{i}",
            "acertos": int(acertos[i]), "theta": float(thetas[i]),
            "se": float(ses[i]), "acertos_faceis": 0, "acertos_medios": 0,
            "acertos_dificeis": 0, "acert_dif": 0.0,
        })
    res = ancorar_notas(thetas_list, tab, area)
    return np.array([r["nota_ancorada"] for r in res])


def aplicar_sigma_mu(thetas: np.ndarray, sigma: float, mu: float) -> np.ndarray:
    return np.clip(sigma * thetas + mu, 0.0, 1000.0)


def processar(area: str, ano: int, resp_mat: np.ndarray, params_of: Dict,
              tabelas: Dict[int, TabelaReferenciaTRI],
              excluir_idx: List[int] = None) -> Dict:
    params = params_para_dicts(params_of[str(ano)][area])
    if excluir_idx:
        keep = [j for j in range(len(params)) if j not in excluir_idx]
        params = [params[j] for j in keep]
        resp = resp_mat[:, keep]
    else:
        resp = resp_mat

    primi = computar_primi(resp, params)
    thetas, ses = estimar_thetas(resp, params)
    acertos = resp.sum(axis=1)

    if ano == 2024:
        sigma = SIGMA_CN_2024 if area == "CN" else SIGMA_MT_2024
        mu = MU_CN_2024 if area == "CN" else MU_MT_2024
        notas = aplicar_sigma_mu(thetas, sigma, mu)
    else:
        notas = ancorar_tabela(area, thetas, acertos, ses, tabelas[ano])

    return {
        "thetas": thetas, "ses": ses, "acertos": acertos, "notas": notas,
        "n_itens": len(params), "primi": primi, "params": params,
    }


# ────────────────────────────────────────────────────────────────────────────
# GRÁFICOS estilo PRIMI
# ────────────────────────────────────────────────────────────────────────────

def desenhar_zonas_pedagogicas(ax, x_min=0, x_max=1000, eixo="y"):
    """Sobrepõe as 5 faixas pedagógicas como bandas horizontais (eixo='y') ou
    verticais (eixo='x')."""
    faixas = [
        (650, 1000, "AVANÇADO"),
        (580, 650, "PROFICIENTE"),
        (480, 580, "BÁSICO"),
        (380, 480, "ELEMENTAR"),
        (0, 380, "INICIANTE"),
    ]
    for ymin, ymax, nv in faixas:
        if eixo == "y":
            ax.axhspan(ymin, ymax, alpha=0.10, color=NIVEIS_CORES[nv], zorder=0)
            ax.text(x_max * 0.98, (ymin + ymax) / 2, nv,
                    fontsize=8, ha="right", va="center",
                    color=NIVEIS_CORES[nv], fontweight="bold", alpha=0.8)
        else:
            ax.axvspan(ymin, ymax, alpha=0.10, color=NIVEIS_CORES[nv], zorder=0)


def plot_acertos_x_nota_com_banda(ax, area: str, ano: int,
                                   acertos: np.ndarray, notas: np.ndarray,
                                   tab: TabelaReferenciaTRI, n_itens: int):
    """Painel PRIMI estilo: acertos × nota V4, com banda INEP (tri_min/med/max do ano)."""
    # Banda de referência INEP por acerto
    ac_axis = np.arange(0, n_itens + 1)
    tri_min = np.array([tab.obter(area, a)["tri_min"] for a in ac_axis])
    tri_med = np.array([tab.obter(area, a)["tri_med"] for a in ac_axis])
    tri_max = np.array([tab.obter(area, a)["tri_max"] for a in ac_axis])

    ax.fill_between(ac_axis, tri_min, tri_max, alpha=0.20,
                    color="#888", label=f"Banda INEP {ano} (tri_min..tri_max)", zorder=1)
    ax.plot(ac_axis, tri_med, "--", color="#444", linewidth=1.5,
            label=f"Mediana INEP {ano}", zorder=2)

    # Zonas pedagógicas no eixo Y
    desenhar_zonas_pedagogicas(ax, x_min=0, x_max=n_itens, eixo="y")

    # Pontos dos alunos
    cores = [NIVEIS_CORES[nivel_pedagogico(n)] for n in notas]
    ax.scatter(acertos, notas, s=22, c=cores, edgecolor="white",
               linewidth=0.5, alpha=0.85, zorder=3)

    ax.set_xlim(-0.5, n_itens + 0.5)
    ax.set_ylim(200, 1000)
    ax.set_xlabel(f"Acertos {area} (0 a {n_itens})")
    ax.set_ylabel(f"Nota {area} V4 INEP {ano}")
    ax.set_title(f"{area} {ano} — acertos × nota com banda INEP",
                 fontsize=12, fontweight="bold")
    ax.legend(loc="upper left", fontsize=9, framealpha=0.85)
    ax.grid(True, alpha=0.25, zorder=0)


def plot_acert_dif_x_nota(ax, area: str, ano: int,
                           acert_dif: np.ndarray, notas: np.ndarray,
                           n_dif: int):
    """Painel PRIMI: acert_dif (% itens difíceis acertados) × nota V4."""
    desenhar_zonas_pedagogicas(ax, x_min=0, x_max=1, eixo="y")

    # Pontos dos alunos
    cores = [NIVEIS_CORES[nivel_pedagogico(n)] for n in notas]
    # Adicionar jitter ao acert_dif pra evitar empilhamento em valores discretos
    rng = np.random.default_rng(42)
    jitter = rng.uniform(-0.005, 0.005, size=len(acert_dif))
    ax.scatter(acert_dif + jitter, notas, s=22, c=cores, edgecolor="white",
               linewidth=0.5, alpha=0.85, zorder=3)

    # Linha esperada: alunos com alto acert_dif devem ter notas altas
    # Regressão linear simples como referência teórica
    if len(acert_dif) > 1:
        coef = np.polyfit(acert_dif, notas, 1)
        xr = np.linspace(0, 1, 50)
        ax.plot(xr, np.polyval(coef, xr), "--", color="#222",
                linewidth=1.2, alpha=0.7,
                label=f"Tendência local (β={coef[0]:.0f})", zorder=2)

    ax.set_xlim(-0.05, 1.05)
    ax.set_ylim(200, 1000)
    ax.set_xlabel(f"acert_dif: % itens DIFÍCEIS acertados (b>1.5)\n"
                  f"{n_dif} itens INEP {ano} classificados como difíceis")
    ax.set_ylabel(f"Nota {area} V4 INEP {ano}")
    ax.set_title(f"{area} {ano} — PRIMI: acert_dif × nota",
                 fontsize=12, fontweight="bold")
    ax.legend(loc="upper left", fontsize=9, framealpha=0.85)
    ax.grid(True, alpha=0.25, zorder=0)


def gerar_grafico_primi(df: pd.DataFrame, ano: int, png_path: Path,
                         resultados: Dict, tabelas: Dict):
    fig, axes = plt.subplots(2, 2, figsize=(18, 14), dpi=150)

    r_cn = resultados[("CN", ano)]
    r_mt = resultados[("MT", ano)]

    # Linha 1: acertos × nota com banda INEP
    plot_acertos_x_nota_com_banda(
        axes[0, 0], "CN", ano,
        r_cn["acertos"], r_cn["notas"], tabelas[ano], r_cn["n_itens"],
    )
    plot_acertos_x_nota_com_banda(
        axes[0, 1], "MT", ano,
        r_mt["acertos"], r_mt["notas"], tabelas[ano], r_mt["n_itens"],
    )

    # Linha 2: acert_dif PRIMI × nota
    plot_acert_dif_x_nota(
        axes[1, 0], "CN", ano,
        r_cn["primi"]["acert_dif"], r_cn["notas"], r_cn["primi"]["n_dificeis"],
    )
    plot_acert_dif_x_nota(
        axes[1, 1], "MT", ano,
        r_mt["primi"]["acert_dif"], r_mt["notas"], r_mt["primi"]["n_dificeis"],
    )

    plt.suptitle(
        f"FREDÃO Sim2 — V4 INEP escala {ano} — análise PRIMI\n"
        "Linha 1: acertos × nota com banda de referência INEP por ano  |  "
        "Linha 2: % itens DIFÍCEIS acertados (PRIMI) × nota",
        fontsize=14, fontweight="bold", y=1.00,
    )
    plt.tight_layout()
    plt.savefig(png_path, dpi=200, bbox_inches="tight")
    plt.close()


# ────────────────────────────────────────────────────────────────────────────
# PIPELINE
# ────────────────────────────────────────────────────────────────────────────

def main():
    t0 = time.time()
    print("FREDÃO Sim2 — V4 INEP + gráficos PRIMI")
    print("=" * 70)

    ids_cn, mat_cn = carregar(CN_PATH)
    ids_mt, mat_mt = carregar(MT_PATH)
    if ids_cn != ids_mt:
        raise RuntimeError("IDs CN ≠ MT")
    ids = ids_cn
    print(f"  Alunos: {len(ids)}")

    params_of = carregar_params()
    tabelas = TabelaReferenciaTRI.carregar_multiano()

    resultados = {}
    for ano in [2017, 2024]:
        resultados[("CN", ano)] = processar(
            "CN", ano, mat_cn, params_of, tabelas, excluir_idx=None,
        )
        resultados[("MT", ano)] = processar(
            "MT", ano, mat_mt, params_of, tabelas, excluir_idx=[MT_ANULADA_IDX],
        )
        primi_cn = resultados[("CN", ano)]["primi"]
        primi_mt = resultados[("MT", ano)]["primi"]
        print(f"\n  {ano}:")
        print(f"    CN: {primi_cn['n_faceis']} fáceis + {primi_cn['n_medios']} médios + "
              f"{primi_cn['n_dificeis']} difíceis (b>1.5)")
        print(f"    MT: {primi_mt['n_faceis']} fáceis + {primi_mt['n_medios']} médios + "
              f"{primi_mt['n_dificeis']} difíceis (b>1.5)")
        print(f"    CN nota μ={resultados[('CN', ano)]['notas'].mean():.1f}, "
              f"acert_dif μ={primi_cn['acert_dif'].mean():.3f}")
        print(f"    MT nota μ={resultados[('MT', ano)]['notas'].mean():.1f}, "
              f"acert_dif μ={primi_mt['acert_dif'].mean():.3f}")

    # Build DataFrame
    rows = []
    for i, aluno_id in enumerate(ids):
        row = {"ID": aluno_id}
        for area in ["CN", "MT"]:
            row[f"ACERTOS {area}"] = int(resultados[(area, 2024)]["acertos"][i])
        for ano in [2017, 2024]:
            for area in ["CN", "MT"]:
                r = resultados[(area, ano)]
                row[f"θ {area} {ano}"] = round(float(r["thetas"][i]), 3)
                row[f"NOTA {area} {ano}"] = round(float(r["notas"][i]), 1)
                row[f"NÍVEL {area} {ano}"] = nivel_pedagogico(float(r["notas"][i]))
                row[f"ac_dif {area} {ano}"] = round(float(r["primi"]["acert_dif"][i]), 3)
                row[f"ac_dif_n {area} {ano}"] = int(r["primi"]["acertos_dificeis"][i])
            ncn = float(resultados[("CN", ano)]["notas"][i])
            nmt = float(resultados[("MT", ano)]["notas"][i])
            med = (ncn + nmt) / 2
            row[f"MÉDIA {ano}"] = round(med, 1)
            row[f"NÍVEL MÉDIO {ano}"] = nivel_pedagogico(med)
        row["DIFF MÉDIA 24-17"] = round(row["MÉDIA 2024"] - row["MÉDIA 2017"], 1)
        rows.append(row)

    df = pd.DataFrame(rows)
    df["RANK MÉDIA 2024"] = df["MÉDIA 2024"].rank(ascending=False, method="min").astype(int)
    df = df.sort_values("RANK MÉDIA 2024").reset_index(drop=True)

    # Gráficos
    ts = time.strftime("%Y%m%d_%H%M%S")
    png_2017 = OUT_DIR / f"xtri_FREDAO_sim2_V4_PRIMI_2017_{ts}.png"
    png_2024 = OUT_DIR / f"xtri_FREDAO_sim2_V4_PRIMI_2024_{ts}.png"
    gerar_grafico_primi(df, 2017, png_2017, resultados, tabelas)
    gerar_grafico_primi(df, 2024, png_2024, resultados, tabelas)
    print(f"\n  PNG 2017: {png_2017.name}")
    print(f"  PNG 2024: {png_2024.name}")

    # Resumo
    def por_nivel(df: pd.DataFrame, ano: int) -> pd.DataFrame:
        col_n = f"NÍVEL MÉDIO {ano}"
        ordem = ["AVANÇADO", "PROFICIENTE", "BÁSICO", "ELEMENTAR", "INICIANTE"]
        faixas = {"AVANÇADO": "≥650", "PROFICIENTE": "≥580", "BÁSICO": "≥480",
                  "ELEMENTAR": "≥380", "INICIANTE": "<380"}
        rows_n = []
        total = len(df)
        for nv in ordem:
            sub = df[df[col_n] == nv]
            n = len(sub)
            if n == 0:
                rows_n.append({"NÍVEL": nv, "FAIXA": faixas[nv], "N": 0, "%": 0.0,
                               "CN MED": None, "MT MED": None, "MÉDIA MED": None,
                               "ac_dif CN MED": None, "ac_dif MT MED": None})
                continue
            rows_n.append({
                "NÍVEL": nv, "FAIXA": faixas[nv], "N": n,
                "%": round(100 * n / total, 1),
                "CN MED": round(sub[f"NOTA CN {ano}"].median(), 1),
                "MT MED": round(sub[f"NOTA MT {ano}"].median(), 1),
                "MÉDIA MED": round(sub[f"MÉDIA {ano}"].median(), 1),
                "ac_dif CN MED": round(sub[f"ac_dif CN {ano}"].median(), 3),
                "ac_dif MT MED": round(sub[f"ac_dif MT {ano}"].median(), 3),
            })
        return pd.DataFrame(rows_n)

    df_niv_2017 = por_nivel(df, 2017)
    df_niv_2024 = por_nivel(df, 2024)

    elapsed = time.time() - t0
    n_dif_cn_2024 = resultados[("CN", 2024)]["primi"]["n_dificeis"]
    n_dif_mt_2024 = resultados[("MT", 2024)]["primi"]["n_dificeis"]
    n_dif_cn_2017 = resultados[("CN", 2017)]["primi"]["n_dificeis"]
    n_dif_mt_2017 = resultados[("MT", 2017)]["primi"]["n_dificeis"]

    resumo = pd.DataFrame([
        ("Simulado", "Fredão — Simulado 2"),
        ("Análise", "V4 INEP OFICIAL + métricas PRIMI"),
        ("Alunos", len(ids)),
        ("Tempo", f"{elapsed:.1f}s"),
        ("── Classificação PRIMI (b>1.5 = difícil) ──", ""),
        ("CN 2017 — itens difíceis", f"{n_dif_cn_2017}/45"),
        ("CN 2024 — itens difíceis", f"{n_dif_cn_2024}/45"),
        ("MT 2017 — itens difíceis", f"{n_dif_mt_2017}/44"),
        ("MT 2024 — itens difíceis", f"{n_dif_mt_2024}/44"),
        ("── Escala 2017 V4 híbrido ──", ""),
        ("CN nota μ ± σ", f"{df['NOTA CN 2017'].mean():.1f} ± {df['NOTA CN 2017'].std():.1f}"),
        ("MT nota μ ± σ", f"{df['NOTA MT 2017'].mean():.1f} ± {df['NOTA MT 2017'].std():.1f}"),
        ("MÉDIA μ ± σ", f"{df['MÉDIA 2017'].mean():.1f} ± {df['MÉDIA 2017'].std():.1f}"),
        ("ac_dif CN μ", f"{df['ac_dif CN 2017'].mean():.3f}"),
        ("ac_dif MT μ", f"{df['ac_dif MT 2017'].mean():.3f}"),
        ("── Escala 2024 V4 puro ──", ""),
        ("CN nota μ ± σ", f"{df['NOTA CN 2024'].mean():.1f} ± {df['NOTA CN 2024'].std():.1f}"),
        ("MT nota μ ± σ", f"{df['NOTA MT 2024'].mean():.1f} ± {df['NOTA MT 2024'].std():.1f}"),
        ("MÉDIA μ ± σ", f"{df['MÉDIA 2024'].mean():.1f} ± {df['MÉDIA 2024'].std():.1f}"),
        ("ac_dif CN μ", f"{df['ac_dif CN 2024'].mean():.3f}"),
        ("ac_dif MT μ", f"{df['ac_dif MT 2024'].mean():.3f}"),
        ("── Comparação ──", ""),
        ("DIFF MÉDIA 2024-2017 μ", f"{df['DIFF MÉDIA 24-17'].mean():+.1f}"),
    ], columns=["Métrica", "Valor"])

    # Excel
    xlsx_path = OUT_DIR / f"xtri_FREDAO_sim2_V4_PRIMI_{ts}.xlsx"
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        resumo.to_excel(writer, sheet_name="Resumo Executivo", index=False)
        df.to_excel(writer, sheet_name="Alunos", index=False)
        df_niv_2017.to_excel(writer, sheet_name="Por Nivel 2017", index=False)
        df_niv_2024.to_excel(writer, sheet_name="Por Nivel 2024", index=False)
        wb = writer.book
        from openpyxl.drawing.image import Image as XLImage
        ws_g17 = wb.create_sheet("Grafico 2017 PRIMI")
        ws_g17.add_image(XLImage(str(png_2017)), "B2")
        ws_g24 = wb.create_sheet("Grafico 2024 PRIMI")
        ws_g24.add_image(XLImage(str(png_2024)), "B2")

    wb = openpyxl.load_workbook(xlsx_path)
    from openpyxl.styles import Font, Alignment
    bold = Font(bold=True)
    centro = Alignment(horizontal="center")
    for sheet in ["Resumo Executivo", "Alunos", "Por Nivel 2017", "Por Nivel 2024"]:
        ws = wb[sheet]
        for cell in ws[1]:
            cell.font = bold
            cell.alignment = centro
        for col in ws.columns:
            max_len = max((len(str(c.value)) if c.value is not None else 0) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 30)
    wb.save(xlsx_path)
    print(f"  XLSX: {xlsx_path.name}")
    print(f"\nTOTAL: {elapsed:.1f}s")

    print("\n── Por Nível MÉDIO 2017 ──")
    print(df_niv_2017.to_string(index=False))
    print("\n── Por Nível MÉDIO 2024 ──")
    print(df_niv_2024.to_string(index=False))


if __name__ == "__main__":
    main()
