"""Fredão Sim2 — CN ancorada em DUAS escalas pedagógicas (2017 e 2024).

Ancoragem por tabela de referência multianos (tri_min/tri_med/tri_max).
CN tem 45 itens, zero anuladas, zero NaN.
Gera 1 Excel + 2 PNGs separados.
"""

from __future__ import annotations

import sys
import time
from pathlib import Path
from typing import Dict, List, Tuple

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from app.irt.scoring import eap_theta  # noqa: E402
from app.irt.transform import ancorar_notas  # noqa: E402
from app.irt.unidim_em import calibrar_itens_mml_em  # noqa: E402
from app.tabela_referencia import TabelaReferenciaTRI  # noqa: E402

DROPBOX = Path("/Users/home/Library/CloudStorage/Dropbox/Fredao Simulados")
CN_PATH = DROPBOX / "Natureza - Simulado 2.xlsx"

ANOS = [2017, 2024]

OUT_DIR = Path("/Users/home/Downloads/XTRI Engine")
OUT_DIR.mkdir(parents=True, exist_ok=True)


def carregar_cn() -> Tuple[List[str], np.ndarray]:
    wb = openpyxl.load_workbook(CN_PATH, data_only=True)
    ws = wb.worksheets[0]
    ids = []
    rows = []
    for r in range(2, ws.max_row + 1):
        aluno = ws.cell(r, 1).value
        if not aluno:
            continue
        ids.append(str(aluno))
        rows.append([int(ws.cell(r, c).value or 0) for c in range(2, 47)])
    return ids, np.array(rows, dtype=int)


def nivel_pedagogico(nota: float) -> str:
    if nota >= 650:
        return "AVANÇADO"
    if nota >= 580:
        return "PROFICIENTE"
    if nota >= 480:
        return "BÁSICO"
    if nota >= 380:
        return "ELEMENTAR"
    return "INICIANTE"


def gerar_grafico_ano(df: pd.DataFrame, ano: int, png_path: Path):
    fig, axes = plt.subplots(1, 2, figsize=(15, 6), dpi=150)
    bins = np.linspace(0, 1000, 51)

    notas_col = f"NOTA CN {ano}"
    ax = axes[0]
    ax.hist(df[notas_col], bins=bins, color="#2E86AB", edgecolor="white", alpha=0.85)
    mu, sd = df[notas_col].mean(), df[notas_col].std()
    ax.set_title(f"CN escala {ano} — μ={mu:.1f}  σ={sd:.1f}",
                 fontsize=14, fontweight="bold")
    ax.set_xlabel(f"Nota CN (escala pedagógica {ano})")
    ax.set_ylabel("Nº alunos")
    for lim, lab in [(650, "AV"), (580, "PR"), (480, "BA"), (380, "EL")]:
        ax.axvline(lim, color="gray", linestyle="--", alpha=0.5)
        ax.text(lim, ax.get_ylim()[1] * 0.95, lab, fontsize=9,
                ha="center", color="gray")

    ax2 = axes[1]
    ax2.scatter(df["ACERTOS CN"], df[notas_col], s=18, alpha=0.55,
                color="#264653", edgecolor="white", linewidth=0.4)
    ax2.set_xlabel("Acertos CN (0-45)")
    ax2.set_ylabel(f"Nota CN escala {ano}")
    ax2.set_title(f"Curva acertos × nota — escala {ano}",
                  fontsize=14, fontweight="bold")
    ax2.grid(True, alpha=0.3)

    plt.suptitle(f"FREDÃO Simulado 2 — Ciências da Natureza — ancoragem pedagógica {ano}",
                 fontsize=15, fontweight="bold", y=1.02)
    plt.tight_layout()
    plt.savefig(png_path, dpi=200, bbox_inches="tight")
    plt.close()


def main():
    t0 = time.time()
    print("FREDÃO Sim2 — CN em duas escalas (2017 e 2024)")
    print("=" * 70)

    ids, mat = carregar_cn()
    print(f"  Alunos: {len(ids)} | itens: {mat.shape[1]} (sem anuladas, sem NaN)")

    print(f"  Calibrando MML/EM...")
    t1 = time.time()
    params = calibrar_itens_mml_em(mat, c=0.20, n_quad=61)
    elapsed = time.time() - t1
    a_vals = [p["a"] for p in params]
    b_vals = [p["b"] for p in params]
    print(f"  Calibrado em {elapsed:.1f}s | "
          f"a μ={np.mean(a_vals):.2f}±{np.std(a_vals):.2f}, "
          f"b μ={np.mean(b_vals):.2f}±{np.std(b_vals):.2f}")

    thetas_arr = np.zeros(len(ids))
    ses_arr = np.zeros(len(ids))
    acertos_arr = mat.sum(axis=1)

    thetas_list = []
    for i, row in enumerate(mat):
        t, se = eap_theta(row, params)
        thetas_arr[i] = t
        ses_arr[i] = se
        thetas_list.append({
            "aluno_id": ids[i],
            "nome": ids[i],
            "acertos": int(acertos_arr[i]),
            "theta": float(t),
            "se": float(se),
            "acertos_faceis": 0,
            "acertos_medios": 0,
            "acertos_dificeis": 0,
            "acert_dif": 0.0,
        })
    print(f"  θ μ={thetas_arr.mean():.3f}, σ={thetas_arr.std():.3f}")

    tabelas = TabelaReferenciaTRI.carregar_multiano()
    resultados_por_ano: Dict[int, List[Dict]] = {}
    for ano in ANOS:
        if ano not in tabelas:
            raise RuntimeError(f"Ano {ano} não encontrado na tabela multianos")
        tab = tabelas[ano]
        resultados = ancorar_notas([dict(t) for t in thetas_list], tab, "CN")
        resultados_por_ano[ano] = resultados
        notas = [r["nota_ancorada"] for r in resultados]
        print(f"  Escala {ano}: μ={np.mean(notas):.1f} σ={np.std(notas):.1f} "
              f"min={min(notas):.0f} max={max(notas):.0f}")

    rows = []
    for i, aluno_id in enumerate(ids):
        row = {
            "ID": aluno_id,
            "ACERTOS CN": int(acertos_arr[i]),
            "θ CN": round(float(thetas_arr[i]), 3),
            "SE θ CN": round(float(ses_arr[i]), 3),
        }
        for ano in ANOS:
            nota = float(resultados_por_ano[ano][i]["nota_ancorada"])
            row[f"NOTA CN {ano}"] = round(nota, 1)
            row[f"NÍVEL CN {ano}"] = nivel_pedagogico(nota)
        row["DIFF 2024-2017"] = round(row["NOTA CN 2024"] - row["NOTA CN 2017"], 1)
        flags = []
        if abs(thetas_arr[i]) > 2.5:
            flags.append(f"θ={thetas_arr[i]:+.2f}")
        if ses_arr[i] > 0.6:
            flags.append(f"SE={ses_arr[i]:.2f}")
        row["ATÍPICO"] = "; ".join(flags)
        rows.append(row)

    df = pd.DataFrame(rows)
    df["RANK"] = df["NOTA CN 2024"].rank(ascending=False, method="min").astype(int)
    df = df.sort_values("RANK").reset_index(drop=True)

    def por_nivel(df: pd.DataFrame, ano: int) -> pd.DataFrame:
        col_n = f"NÍVEL CN {ano}"
        col_v = f"NOTA CN {ano}"
        ordem = ["AVANÇADO", "PROFICIENTE", "BÁSICO", "ELEMENTAR", "INICIANTE"]
        faixas = {"AVANÇADO": "≥650", "PROFICIENTE": "≥580", "BÁSICO": "≥480",
                  "ELEMENTAR": "≥380", "INICIANTE": "<380"}
        rows_n = []
        total = len(df)
        for nv in ordem:
            sub = df[df[col_n] == nv]
            n = len(sub)
            if n == 0:
                rows_n.append({"NÍVEL": nv, "FAIXA": faixas[nv], "N": 0, "%": 0.0,
                               "NOTA MED": None, "ACERTOS MED": None})
                continue
            rows_n.append({
                "NÍVEL": nv,
                "FAIXA": faixas[nv],
                "N": n,
                "%": round(100 * n / total, 1),
                "NOTA MED": round(sub[col_v].median(), 1),
                "ACERTOS MED": int(sub["ACERTOS CN"].median()),
            })
        return pd.DataFrame(rows_n)

    df_niv_2017 = por_nivel(df, 2017)
    df_niv_2024 = por_nivel(df, 2024)

    ord_full = ["AVANÇADO", "PROFICIENTE", "BÁSICO", "ELEMENTAR", "INICIANTE"]
    comparacao = pd.crosstab(df["NÍVEL CN 2017"], df["NÍVEL CN 2024"],
                              rownames=["Nível 2017 ↓"], colnames=["Nível 2024 →"],
                              dropna=False)
    comparacao = comparacao.reindex(index=ord_full, columns=ord_full, fill_value=0)

    elapsed_total = time.time() - t0
    n_atipicos = int((df["ATÍPICO"] != "").sum())
    resumo = pd.DataFrame([
        ("Simulado", "Fredão — Simulado 2"),
        ("Área", "Ciências da Natureza (CN)"),
        ("Alunos processados", len(ids)),
        ("Itens", mat.shape[1]),
        ("Anuladas", "Nenhuma"),
        ("Motor", "V3 cientifico — MML/EM 2PL, c=0.20 fixo"),
        ("Método de ancoragem", "Tabela de referência (tri_min/tri_med/tri_max)"),
        ("Anos de referência", "2017 e 2024"),
        ("Tempo total", f"{elapsed_total:.1f}s"),
        ("", ""),
        ("── Escala 2017 ──", ""),
        ("Nota μ", f"{df['NOTA CN 2017'].mean():.1f}"),
        ("Nota σ", f"{df['NOTA CN 2017'].std():.1f}"),
        ("Nota min/max", f"{df['NOTA CN 2017'].min():.0f} / {df['NOTA CN 2017'].max():.0f}"),
        ("", ""),
        ("── Escala 2024 ──", ""),
        ("Nota μ", f"{df['NOTA CN 2024'].mean():.1f}"),
        ("Nota σ", f"{df['NOTA CN 2024'].std():.1f}"),
        ("Nota min/max", f"{df['NOTA CN 2024'].min():.0f} / {df['NOTA CN 2024'].max():.0f}"),
        ("", ""),
        ("── Diferença ──", ""),
        ("DIFF 2024-2017 μ", f"{df['DIFF 2024-2017'].mean():+.1f}"),
        ("DIFF 2024-2017 σ", f"{df['DIFF 2024-2017'].std():.1f}"),
        ("DIFF 2024-2017 min/max", f"{df['DIFF 2024-2017'].min():+.1f} / {df['DIFF 2024-2017'].max():+.1f}"),
        ("Atípicos (|θ|>2.5 ou SE>0.6)", n_atipicos),
    ], columns=["Métrica", "Valor"])

    df_atip = df[df["ATÍPICO"] != ""].copy().reset_index(drop=True)

    ts = time.strftime("%Y%m%d_%H%M%S")
    png_2017 = OUT_DIR / f"xtri_FREDAO_sim2_CN_escala2017_{ts}.png"
    png_2024 = OUT_DIR / f"xtri_FREDAO_sim2_CN_escala2024_{ts}.png"
    gerar_grafico_ano(df, 2017, png_2017)
    gerar_grafico_ano(df, 2024, png_2024)
    print(f"  PNG 2017: {png_2017}")
    print(f"  PNG 2024: {png_2024}")

    xlsx_path = OUT_DIR / f"xtri_FREDAO_sim2_CN_2anos_{ts}.xlsx"
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        resumo.to_excel(writer, sheet_name="Resumo Executivo", index=False)
        df.to_excel(writer, sheet_name="Alunos", index=False)
        df_niv_2017.to_excel(writer, sheet_name="Por Nivel 2017", index=False)
        df_niv_2024.to_excel(writer, sheet_name="Por Nivel 2024", index=False)
        comparacao.to_excel(writer, sheet_name="Migracao 2017-2024")
        df_atip.to_excel(writer, sheet_name="Atipicos - Revisar", index=False)
        wb = writer.book
        from openpyxl.drawing.image import Image as XLImage
        ws_g17 = wb.create_sheet("Grafico 2017")
        ws_g17.add_image(XLImage(str(png_2017)), "B2")
        ws_g24 = wb.create_sheet("Grafico 2024")
        ws_g24.add_image(XLImage(str(png_2024)), "B2")

    wb = openpyxl.load_workbook(xlsx_path)
    from openpyxl.styles import PatternFill, Font, Alignment
    rosa = PatternFill(start_color="FFE6F0", end_color="FFE6F0", fill_type="solid")
    bold = Font(bold=True)
    centro = Alignment(horizontal="center")

    ws_al = wb["Alunos"]
    headers = [c.value for c in ws_al[1]]
    atip_col = headers.index("ATÍPICO") + 1
    for r in range(2, ws_al.max_row + 1):
        if ws_al.cell(r, atip_col).value:
            for c in range(1, ws_al.max_column + 1):
                ws_al.cell(r, c).fill = rosa

    for sheet in ["Resumo Executivo", "Alunos", "Por Nivel 2017",
                  "Por Nivel 2024", "Migracao 2017-2024", "Atipicos - Revisar"]:
        ws = wb[sheet]
        for cell in ws[1]:
            cell.font = bold
            cell.alignment = centro
        for col in ws.columns:
            max_len = max((len(str(c.value)) if c.value is not None else 0) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 32)

    wb.save(xlsx_path)
    print(f"  XLSX: {xlsx_path}")
    print(f"\nTOTAL: {elapsed_total:.1f}s")

    print("\n=== Resumo escala 2017 ===")
    print(df_niv_2017.to_string(index=False))
    print("\n=== Resumo escala 2024 ===")
    print(df_niv_2024.to_string(index=False))
    print("\n=== Migração de nível ===")
    print(comparacao.to_string())
    print(f"\nAtípicos: {n_atipicos}")
    print(f"\nDIFF 2024-2017: μ={df['DIFF 2024-2017'].mean():+.1f}, "
          f"min={df['DIFF 2024-2017'].min():+.1f}, max={df['DIFF 2024-2017'].max():+.1f}")


if __name__ == "__main__":
    main()
